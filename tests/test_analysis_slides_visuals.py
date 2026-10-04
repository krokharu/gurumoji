"""Native visual layout regressions. Anonymous saved DTOs only."""
import copy
import io
import tempfile
import time
import unittest
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from gurumoji.analysis_core import fingerprint
from gurumoji.services.analysis_slides import AnalysisSlidesService
from gurumoji.services.analysis_slide_projection import build_visual, visible_text
from gurumoji.services.analysis_slide_templates import load_slide_templates
from gurumoji.vault_registry import VaultRegistry
from test_analysis_slides_generation import broad_fixture

ROOT=Path(__file__).resolve().parents[1]

class VisualSlidesTests(unittest.TestCase):
    def setUp(self):
        self.saved=broad_fixture()
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        base=Path(self.temp.name)
        self.vault=VaultRegistry(base/'data'/'library.sqlite3',base/'software')
        self.service=AnalysisSlidesService(export_result=lambda *_:copy.deepcopy(self.saved),vault_factory=lambda:self.vault,templates=load_slide_templates(ROOT))

    def preview(self):
        self.service.save_design('item-1','run-1')
        return self.service.preview('item-1','run-1')

    def test_semantic_layouts_reduce_fragmented_pages_and_keep_major_caveats(self):
        p=self.preview();kinds={s['visual']['kind'] for s in p['slides'] if s.get('visual')}
        self.assertTrue({'hero','method','evidence_chain','statistics','critique','audit'}<=kinds,kinds)
        self.assertLess(len(p['slides']),15)
        self.assertFalse(p['review_required'])
        issue=next(s for s in p['slides'] if (s.get('visual') or {}).get('kind')=='critique')
        text=visible_text(issue['visual'])
        for x in ('反対事例が未保存','独立した事例','未実行','採用・未解決','不足を認める','中心解釈を保留','旧版','版 2'):
            self.assertIn(x,text)
        audit=next(s for s in p['slides'] if (s.get('visual') or {}).get('kind')=='audit')
        self.assertIn(p['snapshot_signature'],audit['notes'])
        self.assertIn('raw_hash',audit['notes'])

    def test_native_chart_table_and_connectors_with_no_external_assets(self):
        from pptx import Presentation
        p=self.preview();data=self.service.presentation('item-1','run-1',expected_snapshot=p['snapshot_signature']).getvalue()
        prs=Presentation(io.BytesIO(data));charts=[];tables=[];connectors=[]
        for slide in prs.slides:
            charts += [s.chart for s in slide.shapes if s.has_chart]
            tables += [s.table for s in slide.shapes if s.has_table]
            connectors += [s for s in slide.shapes if s.shape_type==9]
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left,0);self.assertGreaterEqual(shape.top,0)
                self.assertLessEqual(shape.left+shape.width,prs.slide_width)
                self.assertLessEqual(shape.top+shape.height,prs.slide_height)
        self.assertEqual(len(charts),1);self.assertTrue(tables);self.assertTrue(connectors)
        self.assertEqual(list(charts[0].series[0].values),[50.0])
        with ZipFile(io.BytesIO(data)) as z:
            self.assertTrue(any(n.startswith('ppt/embeddings/') and n.endswith('.xlsx') for n in z.namelist()))
            for n in z.namelist():
                if n.endswith('.xml') or n.endswith('.rels'): ET.fromstring(z.read(n))
                if n.endswith('.rels'): self.assertNotIn(b'TargetMode="External"',z.read(n))
        # Cross-second generation catches embedded workbook timestamps too.
        time.sleep(1.1)
        self.assertEqual(data,self.service.presentation('item-1','run-1',expected_snapshot=p['snapshot_signature']).getvalue())

    def test_zero_and_multilabel_values_are_preserved_not_normalized(self):
        r=next(r for r in self.saved['raw_results'] if r['role']=='statistics')
        r['raw']['rows']=[{'label':'zero','count':0,'denominator':2,'proportion':0},
                          {'label':'both','count':2,'denominator':2,'proportion':1},
                          {'label':'overlap','count':2,'denominator':2,'proportion':1}]
        r['raw_hash']=fingerprint(r['raw']);p=self.preview()
        v=next(s['visual'] for s in p['slides'] if (s.get('visual') or {}).get('kind')=='statistics')
        self.assertEqual(v['values'],[0,100,100]);self.assertEqual(v['counts'],[0,2,2])

    def test_ambiguous_numbers_and_long_values_fall_back_without_truncation(self):
        for value in (None,True,float('nan'),float('inf'),-1):
            saved=copy.deepcopy(self.saved);r=next(r for r in saved['raw_results'] if r['role']=='statistics')
            r['raw']['rows'][0]['count']=value
            self.assertIsNone(build_visual('specialist_perspectives','数量・統計',saved,['result-statistics'],{}))
        self.saved['run']['current_view']['claims'][0]['text']='長い文'*1000
        p=self.preview();s=next(s for s in p['slides'] if s['section']=='major_results')
        self.assertIsNone(s['visual']);self.assertLessEqual(len(s['paragraphs']),10)

    def test_no_agreement_or_label_history_is_invented(self):
        p=self.preview();alltext='\n'.join(visible_text(s.get('visual')) for s in p['slides'])
        self.assertNotIn('全員一致',alltext);self.assertNotIn('ラベル追加',alltext)
        self.assertIn('タスク履歴は\n未記録',alltext)

    def test_unknown_public_shapes_do_not_silently_disappear(self):
        for role,field,value in [('interpretation','limitations',{'text':'重要な留保'}),
                                 ('interpretation','findings',[{'text':'追加の所見'}]),
                                 ('statistics','analysis_unit','発話'),
                                 ('statistics','annotation_version',999)]:
            saved=copy.deepcopy(self.saved);r=next(r for r in saved['raw_results'] if r['role']==role)
            r['raw'][field]=value
            self.assertIsNone(build_visual('specialist_perspectives',role,saved,[r['result_id']],{}))
        self.saved['run']['unresolved_issues'][0]['missing_evidence']='別の保存値'
        self.assertIsNone(build_visual('limitations','限界',self.saved,[],{},['critic','limitations']))
        self.saved=broad_fixture()
        self.assertIsNone(build_visual('limitations','限界',self.saved,[],{},['limitations']))

    def test_extra_issue_caveat_uses_complete_text_fallback(self):
        self.saved['run']['issues'][0]['limitations']=['追加の重大な留保']
        self.saved['run']['unresolved_issues']=copy.deepcopy(self.saved['run']['issues'])
        p=self.preview()
        self.assertFalse(any((s.get('visual') or {}).get('kind') in {'critique','limitations'} for s in p['slides']))
        self.assertIn('追加の重大な留保',''.join(line for s in p['slides'] for line in s['paragraphs']))

    def test_large_critical_caveat_still_blocks_after_visual_projection(self):
        self.saved['run']['issues'][0]['reason']='重大な理由'*1000
        self.saved['run']['unresolved_issues']=copy.deepcopy(self.saved['run']['issues'])
        p=self.preview();self.assertTrue(p['review_required']);self.assertFalse(p['download']['available'])

if __name__=='__main__':unittest.main()
