"""Saved-design-first PPTX acceptance tests using anonymous fixtures only."""
import copy
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from gurumoji.analysis_core import fingerprint
from gurumoji.services.analysis_slides import AnalysisSlidesService, MAX_LINES, MAX_SLIDES
from gurumoji.services.analysis_slide_templates import load_slide_templates
from gurumoji.vault_registry import VaultRegistry
from test_analysis_slides import fixture

ROOT = Path(__file__).resolve().parents[1]


def reviewed_templates():
    templates = load_slide_templates(ROOT)
    if not templates:
        raise unittest.SkipTest("Reviewed Software Vault design/catalog not yet recorded")
    return templates


def broad_fixture():
    saved = fixture()
    saved["run"]["current_view"]["alternatives"] = ["合成データの代替説明。結論は未確定です。"]
    saved["run"]["current_view"]["unresolved"] = ["追加資料は未保存です。"]
    for role in ("interpretation", "statistics", "verification", "critic"):
        raw = {"summary": "合成の保存済み" + role + "結果", "claims": [{"claim_id": role + "-claim", "kind": "interpretation",
               "text": "合成会話の範囲に限った解釈です。", "evidence_ids": ["ev_1"]}], "limitations": ["合成データのみ"]}
        if role == "statistics":
            raw = {"summary": "固定ラベル版2を対象2発話で集計。欠測1発話。", "method_id": "label_frequency",
                   "denominator": 2, "missing_count": 1, "annotation_version": 2,
                   "rows": [{"label": "合成ラベル", "count": 1, "denominator": 2, "proportion": .5, "evidence_ids": ["ev_1"]}],
                   "limitations": ["複数ラベルの合計割合は100%を超える場合があります。"]}
        if role == "critic":
            raw.update(review_status="issues", reviewed_scope="合成観察の根拠", issues=[{
                "issue_key": "missing-evidence", "severity": "high", "reason": "反対事例が未保存",
                "evidence_ids": ["ev_1"], "target_id": "view:run-1", "target_version": 2,
                "missing_evidence": "独立した事例", "proposed_test": "別の合成事例を検討する（未実行）"}])
        saved["raw_results"].append({"run_id": "run-1", "result_id": "result-" + role, "role": role,
             "raw_hash": fingerprint(raw), "raw": raw, "annotation_version": 2, "dataset_version": "input-fixed",
             "validation_status": "valid", "content_status": "unreviewed", "stale": role == "critic"})
    saved["run"]["issues"] = [{**saved["raw_results"][-1]["raw"]["issues"][0], "issue_id": "issue-1",
                               "result_id": "result-critic", "status": "adopted_unresolved", "stale": True}]
    saved["run"]["unresolved_issues"] = copy.deepcopy(saved["run"]["issues"])
    saved["run"]["critique_responses"] = [{"issue_id": "issue-1", "disposition": "adopt", "reason": "不足を認める",
                                          "impact": "中心解釈を保留"}]
    saved["initial"]["snapshot"]["evidence"][0]["source_refs"] = [{"id": "saved-source-1", "url": "https://example.com/synthetic-source"}]
    saved["initial"]["hash"] = fingerprint(saved["initial"]["snapshot"])
    return saved


class SavedSlidesGenerationTests(unittest.TestCase):
    def setUp(self):
        self.templates = reviewed_templates()
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.saved = broad_fixture()
        self.vault = VaultRegistry(self.root / "data" / "library.sqlite3", self.root / "software")
        self.service = AnalysisSlidesService(export_result=lambda *_: copy.deepcopy(self.saved),
                                            vault_factory=lambda: self.vault, templates=self.templates)

    def save(self):
        return self.service.save_design("item-1", "run-1")

    def files(self):
        return {str(p.relative_to(self.root)): (p.read_bytes(), p.stat().st_mtime_ns)
                for p in self.root.rglob("*") if p.is_file()}

    def test_preview_deterministic_bounded_complete_and_read_only(self):
        self.save()
        before = self.files()
        original = copy.deepcopy(self.saved)
        preview = self.service.preview("item-1", "run-1")
        self.assertEqual(preview, self.service.preview("item-1", "run-1"))
        self.assertEqual(self.saved, original)
        self.assertEqual(self.files(), before)
        self.assertLessEqual(len(preview["slides"]), MAX_SLIDES)
        self.assertGreaterEqual(len(preview["slides"]), 10)
        self.assertEqual({s["section"] for s in preview["slides"]}, {s["key"] for s in next(iter(self.templates.values()))["design"]["sections"]})
        for slide in preview["slides"]:
            if not slide.get("visual"):
                self.assertLessEqual(len(slide["paragraphs"]), MAX_LINES)
            self.assertIn(preview["snapshot_signature"], slide["notes"])
            self.assertIn('"annotation_version": "2"', slide["notes"])
        encoded = json.dumps(preview, ensure_ascii=False)
        for value in ("adopted_unresolved", "target_version", "stale", "denominator", "missing_count", "saved-source-1"):
            self.assertIn(value, encoded)

    def test_editable_pptx_reopens_with_provenance_and_no_external_assets(self):
        from pptx import Presentation
        self.save()
        before = self.files()
        preview = self.service.preview("item-1", "run-1")
        stream = self.service.presentation("item-1", "run-1", expected_snapshot=preview["snapshot_signature"])
        self.assertEqual(stream.getvalue(), self.service.presentation("item-1", "run-1", expected_snapshot=preview["snapshot_signature"]).getvalue())
        deck = Presentation(stream)
        self.assertEqual(len(deck.slides), len(preview["slides"]))
        for slide in deck.slides:
            self.assertTrue(all(shape.has_text_frame or shape.has_chart or shape.has_table or shape.shape_type == 9 for shape in slide.shapes))
            self.assertIn(preview["snapshot_signature"], slide.notes_slide.notes_text_frame.text)
            for shape in slide.shapes:
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)
        stream.seek(0)
        with ZipFile(stream) as archive:
            for name in archive.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    ET.fromstring(archive.read(name))
                if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                    self.assertNotIn(b"<p:pic>", archive.read(name))
                if name.endswith(".rels"):
                    self.assertNotIn(b'TargetMode="External"', archive.read(name))
        self.assertEqual(self.files(), before)

    def test_private_and_quarantined_fields_never_enter_any_projection(self):
        for result in self.saved["raw_results"]:
            result["raw"]["chain_of_thought"] = {"summary": "SECRET_REASONING", "evidence_ids": ["SECRET_EVIDENCE"]}
            result["raw_hash"] = fingerprint(result["raw"])
        for status in (None, "quarantined", "failed", "unknown", "received"):
            self.saved["raw_results"].append({"run_id": "run-1", "result_id": "bad-" + str(status), "role": "interpretation",
                 "validation_status": status, "raw": {"summary": "SECRET_QUARANTINED", "limitations": ["SECRET_QUARANTINED"],
                 "claims": [{"text": "SECRET_QUARANTINED", "evidence_ids": ["SECRET_EVIDENCE"]}]}})
        self.save()
        encoded = json.dumps(self.service.preview("item-1", "run-1"))
        self.assertNotIn("SECRET_", encoded)

    def test_missing_legacy_fields_and_long_japanese_are_explicit(self):
        self.saved["raw_results"] = []
        self.saved["run"]["current_view"] = {"summary": "長い合成文" * 3000}
        self.saved["run"]["issues"] = []
        self.saved["run"]["unresolved_issues"] = []
        self.saved["initial"] = {"initial_id": "initial-1", "hash": None, "snapshot": None}
        self.save()
        preview = self.service.preview("item-1", "run-1")
        self.assertLessEqual(len(preview["slides"]), MAX_SLIDES)
        self.assertTrue(any(s["missing"] for s in preview["slides"]))
        self.assertTrue(any(s["truncated"] for s in preview["slides"]))
        self.assertIn("全文ではありません", json.dumps(preview, ensure_ascii=False))

    def test_material_caveat_truncation_blocks_pptx_even_if_adopted_or_stale(self):
        from gurumoji.analysis_core import AnalysisContractError
        for status, stale in (("adopted_unresolved", True), ("defer", False), ("reject", True)):
            self.saved = broad_fixture()
            self.saved["run"]["issues"][0].update(status=status, stale=stale, reason="重大な未解決の合成理由" * 1000)
            self.saved["run"]["unresolved_issues"] = copy.deepcopy(self.saved["run"]["issues"])
            self.save()
            preview = self.service.preview("item-1", "run-1")
            self.assertTrue(preview["review_required"])
            self.assertFalse(preview["download"]["available"])
            with self.assertRaises(AnalysisContractError) as caught:
                self.service.presentation("item-1", "run-1", expected_snapshot=preview["snapshot_signature"])
            self.assertEqual(caught.exception.code, "slide_review_required")

    def test_omitted_high_issues_fail_closed_and_resolved_is_distinct(self):
        self.saved["run"]["issues"] = [{"issue_id": "critical-" + str(i), "severity": "high", "status": "adopted_unresolved",
            "reason": "合成の重大指摘" + str(i), "target_version": 2} for i in range(30)]
        self.save()
        preview = self.service.preview("item-1", "run-1")
        self.assertTrue(preview["review_required"])
        # This deliberately partial test-only design proves the high-issue gate,
        # independently of a section's page truncation guard.
        self.service.templates = copy.deepcopy(self.templates)
        for template in self.service.templates.values():
            template["design"]["sections"] = [{"key": "purpose", "title": "Synthetic gate check"}]
        self.saved["run"]["issues"] = [{"issue_id": "resolved", "severity": "high", "status": "resolved", "reason": "not displayed"}]
        self.save()
        self.assertFalse(self.service.preview("item-1", "run-1")["review_required"])
        self.saved["run"]["issues"][0]["status"] = "adopted_unresolved"
        self.save()
        self.assertTrue(self.service.preview("item-1", "run-1")["review_required"])

    def test_missing_renderer_returns_503_compatible_error_after_preview(self):
        from gurumoji.analysis_core import AnalysisContractError
        self.save()
        with patch("gurumoji.services.analysis_slides.importlib.util.find_spec", return_value=None):
            preview = self.service.preview("item-1", "run-1")
            self.assertFalse(preview["download"]["available"])
            with self.assertRaises(AnalysisContractError) as caught:
                self.service.presentation("item-1", "run-1", expected_snapshot=preview["snapshot_signature"])
            self.assertEqual(caught.exception.code, "slide_renderer_unavailable")


class SlidesFlaskIntegrationTests(unittest.TestCase):
    def test_real_app_saved_run_design_preview_download_without_calls_or_mutation(self):
        import app
        from pptx import Presentation
        from test_analysis_orchestration_integration import OrchestrationIntegrationTests
        reviewed_templates()
        fixture = OrchestrationIntegrationTests("test_real_adapter_loop_full_initial_export_and_no_poll_model_calls")
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        with patch.object(app, "call_orchestration_ai_json", side_effect=fixture.response):
            run_id = fixture.start()
            fixture.service.run(run_id)
        url = fixture.url + "/" + run_id + "/slides"
        with patch.object(app, "call_orchestration_ai_json", side_effect=AssertionError("No additional AI calls")) as ai, \
             patch.object(app, "analysis_orchestration_service", side_effect=AssertionError("No mutable runtime service")):
            catalog = fixture.client.get(url + "/templates")
            self.assertEqual(catalog.status_code, 200, catalog.get_json())
            self.assertEqual(fixture.client.get(url + "/preview").status_code, 409)
            signature = catalog.get_json()["snapshot_signature"]
            saved = fixture.client.post(url + "/design", json={"template": catalog.get_json()["templates"][0]["id"], "expected_snapshot": signature})
            self.assertEqual(saved.status_code, 200, saved.get_json())
            database_before = Path(app.DATABASE_FILE).read_bytes()
            preview = fixture.client.get(url + "/preview").get_json()
            self.assertEqual(preview["snapshot_signature"], signature)
            download = fixture.client.get(url + "/presentation.pptx", query_string={"template": "research", "expected_snapshot": signature})
            self.assertEqual(download.status_code, 200, download.get_json() if download.is_json else "binary")
            self.assertEqual(len(Presentation(io.BytesIO(download.data)).slides), len(preview["slides"]))
            self.assertIn("attachment", download.headers["Content-Disposition"])
            self.assertEqual(Path(app.DATABASE_FILE).read_bytes(), database_before)
            ai.assert_not_called()
            registry = app.analysis_slides_service().vault_factory()
            receipt = saved.get_json()["receipt"]
            note = registry.root("visualization") / receipt["path"]
            note.unlink()
            self.assertEqual(fixture.client.get(url + "/preview").status_code, 409)
            self.assertEqual(fixture.client.get(url + "/presentation.pptx", query_string={"expected_snapshot": signature}).status_code, 409)
            self.assertFalse(note.exists())


if __name__ == "__main__":
    unittest.main()
