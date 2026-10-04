"""Synthetic regressions for export, denominator, timing and inference contracts."""
import copy
import csv
import io
import itertools
import json
import unittest
from unittest.mock import patch

import app
import research_analysis as research
from gurumoji import analysis_insights
from gurumoji.services import group_analysis
from test_research_analysis import fixture_analysis
import test_analysis as analysis_fixtures


def fallback_context(test):
    for name in ("_load_ginza", "_load_sudachi"):
        manager = patch.object(research, name, return_value=(None, "synthetic offline fixture"))
        manager.start()
        test.addCleanup(manager.stop)
    research._RESEARCH_CACHE.clear()


def word(index):
    return "word" + chr(97 + index // 26) + chr(97 + index % 26)


class AnalysisMetricsPureRegressions(unittest.TestCase):
    def setUp(self):
        fallback_context(self)

    def test_preview_does_not_truncate_complete_export_or_contaminate_cache(self):
        analysis = fixture_analysis("full_rows")
        analysis["segments"] = [{**analysis["segments"][0], "text": " ".join(word(i) for i in range(250))}]
        preview = research.enrich_research_analysis(analysis)
        linguistics = preview["research"]["linguistics"]
        self.assertEqual(len(linguistics["term_frequency"]), 200)
        self.assertTrue(preview["research"]["row_limits"]["term_frequency"]["truncated"])
        self.assertEqual(len(linguistics["speaker_term_frequency"][0]["terms"]), 50)
        # A CSV consumer given the preview still receives the full calculated data.
        sources = research.research_csv_sources(preview)
        self.assertEqual(len(sources["term_frequency"]), 250)
        self.assertAlmostEqual(sum(r["term_percent"] for r in sources["term_frequency"]), 100)
        complete = research.enrich_research_analysis(analysis, include_rows=True)
        self.assertEqual(complete["research"]["row_mode"], "complete")
        self.assertEqual(len(complete["research"]["linguistics"]["term_frequency"]), 250)
        self.assertEqual(len(complete["research"]["linguistics"]["speaker_term_frequency"][0]["terms"]), 250)
        self.assertFalse(complete["research"]["row_limits"]["term_frequency"]["truncated"])
        self.assertEqual(len(sources["cooccurrence"]), 500)
        self.assertEqual(sources["cooccurrence"][0]["dataset_total_rows"], 1770)
        self.assertTrue(sources["cooccurrence"][0]["dataset_truncated"])
        self.assertEqual(sources["cooccurrence"][0]["eligible_term_count"], 60)
        self.assertEqual(sources["cooccurrence"][0]["total_term_count"], 250)

    def test_speaker_df_counts_stopword_only_but_not_blank_or_excluded_utterances(self):
        analysis = fixture_analysis("df_scope")
        source = analysis["segments"][0]
        analysis["segments"] = [{**source, "id": sid, "speaker": speaker, "text": text, "excluded": excluded}
                                for sid, speaker, text, excluded in [
                                    ("cat", "A", "cat", False), ("stop", "A", "はい", False),
                                    ("blank", "A", "  ", False), ("excluded", "A", "cat", True),
                                    ("dog", "B", "dog", False), ("stop_b", "C", "はい", False)]]
        result = research.build_research_analysis(analysis)
        speakers = {r["speaker"]: r for r in result["linguistics"]["speaker_term_frequency"]}
        self.assertEqual(speakers["A"]["document_count"], 2)
        self.assertEqual(speakers["A"]["terms"][0]["document_percent"], 50)
        self.assertEqual(speakers["C"]["document_count"], 1)
        self.assertEqual(speakers["C"]["terms"], [])
        comparison = next(r for r in result["content"]["characteristic_terms"] if r["speaker"] == "A")
        self.assertEqual((comparison["total"], comparison["percent"]), (2, 50))
        self.assertEqual(result["linguistics"]["coverage"]["document_count"], 4)

    def test_crosstab_modes_share_exact_kwic_evidence(self):
        analysis = fixture_analysis("matching")
        source = analysis["segments"][0]
        analysis["segments"] = [{**source, "id": sid, "speaker": speaker, "text": text}
                                for sid, speaker, text in [("e1", "A", "education"), ("e2", "A", "cat"), ("e3", "B", "dog")]]
        analysis["config"]["crosstab_terms"] = ["cat"]
        for mode, expected in [("normalized", ["e2"]), ("surface", ["e2"]), ("literal", ["e1", "e2"])]:
            analysis["config"]["crosstab_match_mode"] = mode
            result = research.build_research_analysis(analysis)
            positive = next(row for row in result["statistics"]["crosstabs"]
                            if row.get("term") == "cat" and row["row_id"] == "A" and row["column_value"] == "あり")
            self.assertEqual(positive["segment_ids"], expected)
            self.assertEqual(positive["count"], len(expected))
            self.assertEqual(positive["match_mode"], mode)
            matches = analysis_insights.search_kwic(analysis, result["linguistics"]["morphemes"], "cat", mode=mode)
            self.assertEqual([row["segment_id"] for row in matches["hits"]], expected)
        self.assertEqual(group_analysis.normalize_analysis_config({"crosstab_match_mode": "bad"})["crosstab_match_mode"], "normalized")
        self.assertEqual(group_analysis.normalize_analysis_config({"crosstab_match_mode": []})["crosstab_match_mode"], "normalized")
        default_hits = analysis_insights.search_kwic(analysis, result["linguistics"]["morphemes"], "cat")
        self.assertEqual(default_hits["mode"], "normalized")
        self.assertEqual([r["segment_id"] for r in default_hits["hits"]], ["e2"])

    def test_normalized_mode_is_not_union_of_lemma_surface_and_normalized(self):
        analysis = fixture_analysis("normalization")
        source = analysis["segments"][0]
        analysis["segments"] = [{**source, "id": "s", "text": "Cats"}]
        tokens = [{"segment_id": "s", "normalized": "cat", "lemma": "feline", "surface": "Cats",
                   "is_content": True, "is_stop": False, "begin": 0, "end": 4}]
        for query, mode, count in [("cat", "normalized", 1), ("feline", "normalized", 0),
                                   ("Cats", "surface", 1), ("cats", "surface", 0)]:
            analysis["config"].update(crosstab_terms=[query], crosstab_match_mode=mode)
            rows = research._segment_dataset(analysis, tokens)
            result = research._statistics_analysis(analysis, rows)
            positives = [r for r in result["crosstabs"] if r.get("term") == query and r["column_value"] == "あり"]
            self.assertEqual(sum(r["count"] for r in positives), count)
            self.assertEqual(analysis_insights.search_kwic(analysis, tokens, query, mode=mode)["total"], count)

    def test_p_value_precision_and_exploratory_missingness_contract(self):
        analysis = fixture_analysis("tiny_p")
        source = analysis["segments"][0]
        analysis["segments"] = [{**source, "id": f"s{i}", "speaker": speaker, "duration": value, "end": value}
                                for i, (speaker, value) in enumerate([(s, v) for s, values in [("A", [1, 2, 3] * 10),
                                                                                           ("B", [3, 4, 5] * 10)] for v in values])]
        analysis["segments"].append({**source, "id": "missing", "valid_time": False})
        result = research._statistics_analysis(analysis, research._segment_dataset(analysis, []))
        anova = next(r for r in result["tests"] if r["outcome"] == "duration_seconds" and r["effect_name"] == "eta_squared")
        self.assertGreater(anova["p_value"], 0)
        self.assertLess(anova["p_value"], 1e-12)
        self.assertEqual((anova["n"], anova["missing"], anova["unit"]), (60, 1, "秒"))
        self.assertEqual(anova["p_value_adjustment"], "none")
        self.assertTrue(anova["exploratory"])
        self.assertAlmostEqual(anova["effect_size"], .6)
        self.assertGreater(json.loads(json.dumps(anova))["p_value"], 0)
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=research.RESEARCH_CSV_FIELDS["statistical_tests"])
        writer.writeheader(); writer.writerow(anova)
        self.assertGreater(float(next(csv.DictReader(io.StringIO(stream.getvalue())))["p_value"]), 0)
        self.assertIsNone(result["inference_policy"]["comparison_family"])
        self.assertFalse(result["inference_policy"]["independence_verified"])
        self.assertTrue(all("missing" in r and r["p_value_adjustment"] == "none" for r in result["correlations"]))

    def test_physical_time_is_permutation_invariant_and_does_not_mutate_source(self):
        rows = [{"id": "late", "speaker": "A", "speaker_name": "A", "start": 10, "end": 12},
                {"id": "unknown", "speaker": "C", "speaker_name": "C", "start": None, "end": None},
                {"id": "early", "speaker": "B", "speaker_name": "B", "start": 0, "end": 11}]
        for permutation in itertools.permutations(rows):
            source = copy.deepcopy(list(permutation)); before = copy.deepcopy(source)
            overlaps, truncated = group_analysis._overlap_candidates(source, .2)
            self.assertFalse(truncated)
            self.assertEqual([(r["start"], r["end"], r["seconds"]) for r in overlaps], [(10, 11, 1)])
            self.assertEqual(group_analysis._long_gaps(source, 3), [])
            self.assertEqual(source, before)

    def test_ai_reference_validation_is_not_semantic_validation(self):
        result = analysis_insights.validate_findings({"findings": [{"category": "shared", "title": "全員賛成",
            "text": "全員が値上げに賛成した", "segment_ids": ["a", "b"]}]},
            {"a": {"speaker": "A", "text": "値上げに反対"}, "b": {"speaker": "B", "text": "値上げに反対"}})[0]
        self.assertEqual(result["interpretation_status"], "unverified_candidate")
        self.assertTrue(result["researcher_review_required"])
        self.assertEqual(result["validation_scope"], "reference_ids_only")
        self.assertIn("全員", result["validation_note"])

    def test_outline_does_not_extend_short_recording_to_bin_width(self):
        rows = analysis_insights._timeline_rows({"timeline": [{"topic_id": "t", "topic_label": "topic", "bin_index": 0,
                                                                          "start": 0, "end": 20, "segment_count": 1}]}, [], 300)
        self.assertEqual(rows[0]["end"], 20)


class AnalysisMetricsApiRegressions(unittest.TestCase):
    def setUp(self):
        fallback_context(self)
        self.fixture = analysis_fixtures.AnalysisApiTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.addCleanup(self.fixture.tearDown)
        self.client = self.fixture.client

    def test_normal_export_routes_keep_250_words_and_record_cooccurrence_limit(self):
        from openpyxl import load_workbook
        item_id = self.fixture.create_analysis_item("full_export", [{"id": "words", "start": 0, "end": 10,
            "speaker": "A", "text": " ".join(word(i) for i in range(250))}])
        url = f"/api/library/{item_id}/analysis"
        before = self.client.get(url).get_json()
        response = self.client.put(url, json={"source_revision": before["item"]["revision_count"],
            "analysis_revision": before["item"]["analysis_revision"], "config": {"cooccurrence_min_count": 1}})
        self.assertEqual(response.status_code, 200, response.data)
        preview = self.client.get(url + "?execute=1").get_json()
        self.assertEqual(len(preview["research"]["linguistics"]["term_frequency"]), 200)
        response = self.client.get(url + "/export.json")
        self.assertEqual(response.status_code, 200, response.data[:300])
        payload = response.get_json()
        self.assertEqual(len(payload["research"]["linguistics"]["term_frequency"]), 250)
        csv_response = self.client.get(url + "/export.csv?dataset=term_frequency")
        self.assertEqual(csv_response.status_code, 200)
        csv_rows = list(csv.DictReader(io.StringIO(csv_response.data.decode("utf-8-sig"))))
        self.assertEqual(len(csv_rows), 250)
        self.assertAlmostEqual(sum(float(r["term_percent"]) for r in csv_rows), 100)
        response = self.client.get(url + "/export.xlsx")
        self.assertEqual(response.status_code, 200, response.data[:300])
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["語彙頻度"].max_row - 1, 250)
        self.assertEqual(workbook["共起"].max_row - 1, 500)
        cooccurrence = self.client.get(url + "/export.csv?dataset=cooccurrence")
        row = next(csv.DictReader(io.StringIO(cooccurrence.data.decode("utf-8-sig"))))
        self.assertEqual(row["dataset_total_rows"], "1770")
        self.assertEqual(row["dataset_returned_rows"], "500")
        self.assertEqual(row["dataset_truncated"], "1")

    def test_actual_participant_count_is_separate_from_observed_speakers(self):
        item_id = self.fixture.create_analysis_item("participation_scope", [{"id": "a", "start": 0, "end": 1,
            "speaker": "A", "text": "cat"}])
        url = f"/api/library/{item_id}/analysis"
        prepared = self.client.get(url).get_json()["manual"]["preparation"]
        response = self.client.put(f"/api/library/{item_id}/preparation", json={
            "revision": prepared["revision"], "source_hash": prepared["source_hash"], "participant_count": 3,
            "metadata_sources": "synthetic roster count", "records": {"a": {"speaker_verified": True,
            "role": "participant", "text_status": "transcript_checked"}}})
        self.assertEqual(response.status_code, 200, response.data)
        result = self.client.get(url).get_json()
        balance = result["automatic"]["balance"]
        self.assertEqual(balance["participant_count"], 3)
        self.assertEqual(balance["observed_participant_count"], 1)
        self.assertEqual(result["automatic"]["overview"]["participant_count"], 3)
        self.assertEqual(balance["denominator"], "observed_participant_speakers")
        self.assertIn("観測発言者内", balance["denominator_label"])
        self.assertIsNone(balance["unobserved_participant_count"])
        self.assertFalse(balance["participant_roster_available"])

    def test_unknown_and_missing_time_speakers_do_not_become_zero_second_participants(self):
        item_id = self.fixture.create_analysis_item("unknown_scope", [
            {"id": "a", "start": 0, "end": 1, "speaker": "A", "text": "cat"},
            {"id": "b", "time_unknown": True, "speaker": "B", "text": "dog"},
            {"id": "u", "start": 2, "end": 3, "speaker": "UNKNOWN", "text": "bird"}])
        result = self.client.get(f"/api/library/{item_id}/analysis").get_json()
        balance = result["automatic"]["balance"]
        self.assertIsNone(balance["participant_count"])
        self.assertEqual(balance["observed_participant_count"], 2)
        self.assertEqual(balance["balance_speaker_count"], 1)
        self.assertEqual(balance["unknown_speaker_turn_count"], 1)
        self.assertEqual(balance["missing_time_turn_count"], 1)
        metrics = {r["speaker"]: r for r in result["automatic"]["speaker_metrics"]}
        self.assertFalse(metrics["B"]["included_in_balance"])
        self.assertFalse(metrics["UNKNOWN"]["included_in_balance"])

    def test_unsorted_saved_transcript_retains_ids_order_and_correct_physical_metrics(self):
        item_id = self.fixture.create_analysis_item("time_scope", [
            {"id": "late", "speaker": "A", "start": 10, "end": 12, "text": "cat"},
            {"id": "unknown", "speaker": "C", "time_unknown": True, "text": "bird"},
            {"id": "early", "speaker": "B", "start": 0, "end": 11, "text": "dog"}])
        result = self.client.get(f"/api/library/{item_id}/analysis").get_json()
        self.assertEqual([r["id"] for r in result["segments"]], ["late", "unknown", "early"])
        self.assertEqual(result["segments"][0]["next_segment_id"], "unknown")
        self.assertEqual(result["automatic"]["overlap_candidates"][0]["seconds"], 1)
        self.assertEqual(result["automatic"]["long_gaps"], [])


if __name__ == "__main__":
    unittest.main()
