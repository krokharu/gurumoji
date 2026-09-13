import copy
import csv
import io
import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

import app
import analysis_insights as content
import research_analysis


def fixture():
    segments = [
        {"id": "a1", "speaker": "A", "speaker_name": "参加者A", "text": "価格を価格で比べる", "start": 0, "end": 2},
        {"id": "a2", "speaker": "A", "speaker_name": "参加者A", "text": "良い", "start": 3, "end": 5},
        {"id": "b1", "speaker": "B", "speaker_name": "参加者B", "text": "価格より操作を改善", "start": 6, "end": 8},
        {"id": "x1", "speaker": "B", "speaker_name": "参加者B", "text": "除外の価格", "start": 9, "end": 11, "excluded": True},
        {"id": "empty", "speaker": "A", "text": " ", "start": 12, "end": 13},
    ]
    for row in segments:
        row.setdefault("annotation", {})
        row.setdefault("excluded", False)
        row.setdefault("role", "participant")
    tokens = []
    for sid, surface, term, begin in [
        ("a1", "価格", "価格", 0), ("a1", "価格", "価格", 3),
        ("a2", "良い", "良い", 0), ("b1", "価格", "価格", 0), ("b1", "操作", "操作", 4),
        ("x1", "価格", "価格", 3),
    ]:
        tokens.append({"segment_id": sid, "surface": surface, "normalized": term,
                       "begin": begin, "end": begin + len(surface), "is_content": True,
                       "is_stop": False, "excluded": sid == "x1"})
    return {"segments": segments, "item": {"id": "content", "revision_count": 0, "analysis_revision": 0},
            "config": {}, "automatic": {}, "manual": {}}, tokens


def finding(ids=None, category="themes", text="価格の話題を確認できます。"):
    return {"findings": [{"category": category, "title": "価格の話題", "text": text,
                          "segment_ids": ids if ids is not None else ["a1"]}]}


class ContentAnalysisTests(unittest.TestCase):
    def test_difference_counts_documents_and_includes_tokenless_denominators(self):
        analysis, tokens = fixture()
        # A's second utterance has no content tokens, but remains in the denominator.
        result = content.build_content_analysis(analysis, [t for t in tokens if t["segment_id"] != "a2"])
        price = next(r for r in result["characteristic_terms"] if r["speaker"] == "B" and r["term"] == "価格")
        self.assertEqual((price["count"], price["total"], price["other_count"], price["other_total"]), (1, 1, 1, 2))
        self.assertEqual(price["difference_pp"], 50)
        shared = next(r for r in result["local"] if r.get("term") == "価格")
        self.assertEqual(shared["segment_ids"], ["a1", "b1"])
        self.assertEqual(shared["count"], 2)
        self.assertLessEqual(len(result["local"]), 5)

    def test_fallback_hiragana_fragments_are_searchable_but_not_observations(self):
        analysis, _ = fixture()
        tokens, _ = research_analysis._analyze_with_fallback(analysis["segments"], set())
        result = content.build_content_analysis(analysis, tokens)
        self.assertTrue(any(r.get("term") == "価格" for r in result["local"]))
        self.assertFalse(any(r.get("term") in {"を", "より", "い", "べる"}
                             for r in result["local"] + result["characteristic_terms"]))
        self.assertEqual(content.search_kwic(analysis, tokens, "を", mode="normalized")["total"], 2)

    def test_manual_codes_quotes_and_single_speaker(self):
        analysis, tokens = fixture()
        analysis["segments"] = analysis["segments"][:2]
        analysis["segments"][0]["annotation"] = {"codes": ["price"], "important": True}
        analysis["manual"] = {"codebook": [{"id": "price", "label": "価格"}]}
        result = content.build_content_analysis(analysis, tokens)
        self.assertEqual(result["characteristic_terms"], [])
        self.assertTrue({"codes", "quotes"} <= {r["category"] for r in result["local"]})
        self.assertEqual(content.build_content_analysis({"segments": []}, []), {"local": [], "characteristic_terms": []})

    def test_kwic_multiple_hits_normalization_context_and_paging(self):
        analysis, tokens = fixture()
        result = content.search_kwic(analysis, tokens, "価格", limit=1, offset=1)
        self.assertEqual(result["total"], 3)
        self.assertEqual(result["hits"][0]["begin"], 3)
        self.assertEqual(result["hits"][0]["context_ids"], ["a1", "a2"])
        tokens[2]["normalized"] = "良し"
        normalized = content.search_kwic(analysis, tokens, "良し", mode="normalized")
        self.assertEqual(normalized["hits"][0]["match"], "良い")
        self.assertEqual(content.search_kwic(analysis, tokens, "価格", speaker="B")["total"], 1)
        self.assertEqual(content.search_kwic(analysis, tokens, ".*")["total"], 0)

    def test_kwic_uses_tokens_past_preview_and_preserves_unicode_offsets(self):
        analysis, tokens = fixture()
        analysis["segments"][0]["text"] = "😀" * 50 + "價格" + "後" * 60
        tokens = [dict(tokens[0], normalized="価格", surface="價格", begin=50, end=52)] * 120
        result = content.search_kwic(analysis, tokens, "価格", mode="normalized")
        self.assertEqual(result["total"], 1)
        self.assertEqual(result["hits"][0]["left"], "😀" * 40)
        self.assertEqual(result["hits"][0]["right"], "後" * 40)

    def test_fingerprint_tracks_registry_attributes_and_ignores_ai_and_time(self):
        analysis, _ = fixture()
        before = content.input_fingerprint(analysis)
        analysis["generated_at"] = "tomorrow"
        analysis["insights"] = {"ai": "generated"}
        self.assertEqual(before, content.input_fingerprint(analysis))
        analysis["automatic"]["speaker_metrics"] = [{"speaker": "A", "attributes": {"組織": "new"}}]
        self.assertNotEqual(before, content.input_fingerprint(analysis))

    def test_ai_rejects_missing_unknown_excluded_and_single_speaker_shared_evidence(self):
        analysis, _ = fixture()
        allowed = {s["id"]: s for s in content.included_segments(analysis)}
        for response in [finding([]), finding(["unknown"]), finding(["x1"]), finding(["a1"], "shared"), {"findings": "bad"}]:
            with self.subTest(response=response), self.assertRaises(ValueError):
                content.validate_findings(response, allowed)
        valid = content.validate_findings(finding(["a1", "b1"], "shared"), allowed)
        self.assertEqual(valid[0]["speakers"], ["A", "B"])

    def test_long_ai_input_covers_tail_and_restricts_evidence_at_each_merge(self):
        analysis, _ = fixture()
        analysis["segments"][0]["text"] = "価格" * 16000 + "末尾"
        prompts = []

        def fake_call(system, prompt, schema):
            prompts.append(prompt)
            if "発話記録:\n" in prompt:
                batch = json.loads(prompt.split("発話記録:\n", 1)[1])
                return finding([batch[0]["id"]])
            return finding([schema["properties"]["findings"]["items"]["properties"]["segment_ids"]["items"]["enum"][0]])

        result = content.create_ai_insights(analysis, fake_call, lambda *args: None, lambda: None)
        self.assertGreater(len(prompts), 2)
        self.assertTrue(any("末尾" in p for p in prompts))
        self.assertFalse(any("除外の価格" in p for p in prompts))
        self.assertEqual(result[0]["segment_ids"], ["a1"])

    def test_ai_does_not_return_partial_result_on_failure_or_cancel(self):
        analysis, _ = fixture()
        analysis["segments"][0]["text"] *= 5000
        calls = []

        def failure(*args):
            calls.append(1)
            if len(calls) > 1:
                raise RuntimeError("offline")
            return finding(["E0001"])

        with self.assertRaises(RuntimeError):
            content.create_ai_insights(analysis, failure, lambda *args: None, lambda: None)
        with self.assertRaises(app.InsightCancelled):
            content.create_ai_insights(analysis, failure, lambda *args: None,
                                       lambda: (_ for _ in ()).throw(app.InsightCancelled()))

    def test_ai_retries_once_when_a_local_model_cites_an_adjacent_batch(self):
        analysis, _ = fixture()
        calls = []

        def local_model(*args):
            calls.append(1)
            return finding(["not-in-this-batch"] if len(calls) == 1 else ["E0001"])

        result = content.create_ai_insights(analysis, local_model, lambda *args: None, lambda: None)
        self.assertEqual(len(calls), 2)
        self.assertEqual(result[0]["segment_ids"], ["a1"])


class ContentApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="gurumoji-content-")
        self.addCleanup(self.temp.cleanup)
        self.database_patch = patch.object(app, "DATABASE_FILE", Path(self.temp.name) / "library.sqlite3")
        self.database_patch.start()
        self.addCleanup(self.database_patch.stop)
        app.initialize_library()
        self.ginza = patch.object(research_analysis, "_load_ginza", return_value=(None, "test fallback"))
        self.sudachi = patch.object(research_analysis, "_load_sudachi", return_value=(None, "test fallback"))
        self.ginza.start(); self.sudachi.start()
        self.addCleanup(self.ginza.stop); self.addCleanup(self.sudachi.stop)
        research_analysis._RESEARCH_CACHE.clear()
        self.client = app.app.test_client()
        analysis, _ = fixture()
        app.upsert_library_item(item_id="content", source_name="example.wav", output_dir=Path(self.temp.name) / "output",
                                media_path=None, language="ja", segments=analysis["segments"],
                                speaker_names={"A": "参加者A", "B": "参加者B"}, files=[], outline=None,
                                emotion_analysis=None, write_srt=False, write_json=True)
        self.url = "/api/library/content/analysis"
        data = self.client.get(self.url).get_json()
        config = data["config"]
        self.client.put(self.url, json={"source_revision": 0, "analysis_revision": 0, "config": config,
                                       "annotations": {"x1": {"excluded": True}}})
        self.addCleanup(app.insight_cancel_events.clear)

    def payload(self, request_id="request-content-0001"):
        data = self.client.get(self.url).get_json()
        return {"request_id": request_id, "provider": "openai",
                "source_revision": data["item"]["revision_count"], "analysis_revision": data["item"]["analysis_revision"]}

    def start(self, payload=None):
        with patch.object(app, "load_token_config", return_value=app.TokenConfig(openai_api_key="test-key")), \
             patch.object(app.threading, "Thread") as thread:
            response = self.client.post(self.url + "/insights", json=payload or self.payload())
            args = thread.call_args.kwargs["args"] if thread.called else None
        return response, args

    def run_worker(self, args, response=None, side_effect=None):
        with patch.object(app, "call_ai_json", return_value=response or finding(["E0001"]), side_effect=side_effect):
            app.run_analysis_insight_job(*args)

    def test_local_summary_kwic_full_csv_and_workbook_without_api_calls(self):
        with patch.object(app, "call_ai_json") as ai:
            analysis = self.client.get(self.url).get_json()
            self.assertTrue(analysis["insights"]["local"])
            result = self.client.get(self.url + "/kwic?q=価格&limit=1").get_json()
            self.assertEqual(result["total"], 3)
            self.assertEqual(len(result["hits"]), 1)
            raw = self.client.get(self.url + "/kwic?q=価格&format=csv&limit=1&offset=1")
            self.assertEqual(len(list(csv.DictReader(io.StringIO(raw.data.decode("utf-8-sig"))))), 3)
            ai.assert_not_called()
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(self.client.get(self.url + "/export.xlsx").data), read_only=True)
        self.assertEqual(workbook.sheetnames[1], "分析の見解")
        self.assertIn("話者別特徴語比較", workbook.sheetnames)

    def test_generated_result_persists_and_becomes_stale_on_edit(self):
        response, args = self.start()
        self.assertEqual(response.status_code, 202, response.get_json())
        self.run_worker(args)
        result = self.client.get(self.url + "/insights").get_json()
        self.assertEqual(result["run"]["status"], "completed", result)
        self.assertFalse(result["insights"]["stale"])
        self.assertEqual(result["insights"]["ai"]["evidence"]["a1"]["text"], "価格を価格で比べる")
        app.initialize_library()
        self.assertIsNotNone(self.client.get(self.url).get_json()["insights"]["ai"])
        data = self.client.get(self.url).get_json()
        config = data["config"]
        config["research_question"] = "別の問い"
        self.client.put(self.url, json={"source_revision": 0, "analysis_revision": data["item"]["analysis_revision"], "config": config})
        self.assertTrue(self.client.get(self.url).get_json()["insights"]["stale"])

    def test_insight_job_forwards_the_selected_reasoning_effort(self):
        payload = self.payload()
        payload["ai_efforts"] = {"outline": "low"}
        response, args = self.start(payload)
        self.assertEqual(response.status_code, 202, response.get_json())
        with patch.object(app, "call_ai_json", return_value=finding(["E0001"])) as call:
            app.run_analysis_insight_job(*args)
        self.assertEqual(call.call_args.kwargs["ai_efforts"]["outline"], "low")

    def test_duplicate_request_and_concurrent_job_do_not_repeat_api_work(self):
        payload = self.payload()
        response, args = self.start(payload)
        duplicate, duplicate_args = self.start(payload)
        self.assertEqual(duplicate.status_code, 200)
        self.assertIsNone(duplicate_args)
        busy, busy_args = self.start(self.payload("request-content-0002"))
        self.assertEqual(busy.status_code, 409)
        self.assertIsNone(busy_args)
        self.run_worker(args)
        duplicate, duplicate_args = self.start(payload)
        self.assertEqual(duplicate.status_code, 200)
        self.assertIsNone(duplicate_args)

    def test_status_snapshot_cannot_report_completed_with_old_output(self):
        _, first_args = self.start()
        self.run_worker(first_args)
        _, next_args = self.start(self.payload("request-content-0002"))
        original = app.group_analysis_for_row
        finished = False

        def finish_after_snapshot(row, **kwargs):
            nonlocal finished
            result = original(row, **kwargs)
            if not finished:
                finished = True
                self.run_worker(next_args)
            return result

        with patch.object(app, "group_analysis_for_row", side_effect=finish_after_snapshot):
            status = self.client.get(self.url + "/insights").get_json()
        self.assertEqual(status["run"]["status"], "queued")
        self.assertEqual(status["insights"]["ai"]["request_id"], first_args[0])
        latest = self.client.get(self.url + "/insights").get_json()
        self.assertEqual(latest["run"]["status"], "completed")
        self.assertEqual(latest["insights"]["ai"]["request_id"], next_args[0])

    def test_failure_cancellation_and_source_conflict_preserve_success(self):
        response, args = self.start()
        self.run_worker(args)
        previous = self.client.get(self.url).get_json()["insights"]["ai"]
        _, args = self.start(self.payload("request-content-0002"))
        self.run_worker(args, response=finding(["not-in-transcript"]))
        self.assertEqual(self.client.get(self.url + "/insights").get_json()["run"]["status"], "failed")
        _, args = self.start(self.payload("request-content-0003"))
        self.client.post(self.url + "/insights/cancel", json={"request_id": args[0]})
        self.run_worker(args)
        self.assertEqual(self.client.get(self.url + "/insights").get_json()["run"]["status"], "cancelled")
        _, args = self.start(self.payload("request-content-0004"))
        with app.database_connection() as connection:
            connection.execute("UPDATE library_items SET revision_count=revision_count+1 WHERE id='content'")
        self.run_worker(args)
        status = self.client.get(self.url + "/insights").get_json()
        self.assertEqual(status["run"]["status"], "stale")
        self.assertEqual(status["insights"]["ai"]["request_id"], previous["request_id"])

    def test_interrupted_run_requires_explicit_retry(self):
        response, _ = self.start()
        app.initialize_library()
        status = self.client.get(self.url + "/insights").get_json()
        self.assertEqual(status["run"]["status"], "interrupted")
        self.assertIsNone(status["insights"]["ai"])

    def test_invalid_input_no_keys_and_excluded_evidence(self):
        for query in ["", "?q=test&mode=bad", "?q=test&offset=-1", "?q=test&limit=1000"]:
            self.assertEqual(self.client.get(self.url + "/kwic" + query).status_code, 400)
        self.assertEqual(self.client.get("/api/library/missing/analysis/kwic?q=a").status_code, 404)
        with patch.object(app, "load_token_config", return_value=app.TokenConfig()):
            self.assertEqual(self.client.post(self.url + "/insights", json=self.payload()).status_code, 400)
        payload = self.payload(); payload["source_revision"] = -1
        self.assertEqual(self.client.post(self.url + "/insights", json=payload).status_code, 409)
        _, args = self.start()
        self.run_worker(args, response=finding(["x1"]))
        self.assertEqual(self.client.get(self.url + "/insights").get_json()["run"]["status"], "failed")


if __name__ == "__main__":
    unittest.main()
