"""Saved-run slide projection and editable PPTX, with a saved-design gate.

No model, current-input builder, computation runner, URL/file reader, or publication
callback is imported here. Only ``save_design`` writes, through VaultRegistry's
existing ownership policy. Preview/download consume the same immutable read DTO.
Templates are reviewed application configuration, never request-supplied code.
"""
from __future__ import annotations

import copy
import importlib.util
import io
import json
import math
import re
import unicodedata
from datetime import datetime, timezone
from typing import Any, Callable
from urllib.parse import quote, urlsplit
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from ..analysis_core import AnalysisContractError, canonical, fingerprint
from ..analysis_store import StoreConflict
from .analysis_slide_projection import build_visual, visible_text

SCHEMA_VERSION = "analysis-slides-1"
SECTION_KEYS = frozenset({"purpose", "method", "major_results", "specialist_perspectives",
                          "critic", "core", "limitations", "evidence"})
ROLE_LABELS = {"core": "Core", "handler": "Handler", "interpretation": "会話解釈",
               "statistics": "数量・統計", "verification": "独立検証", "critic": "批判者"}
MAX_SLIDES = 40
MAX_LINES = 10
LINE_UNITS = 72
MAX_NOTES = 64000
MAX_EXPORT_BYTES = 64 * 1024 * 1024
MAX_RESULTS = 500
MAX_REFS = 80
MISSING = "未記録"
DISCLAIMER = "探索的分析・AI下書きを含みます。研究者の確定解釈や因果関係の証明ではありません。"


def _error(message, code="invalid_request", field=""):
    return AnalysisContractError(message, code=code, field=field)


def _mapping(value):
    return value if isinstance(value, dict) else {}


def _rows(value):
    return value if isinstance(value, list) else []


def _text(value, limit=4000):
    """XML-safe plain text. HTML is data; callers must use textContent."""
    if value is None:
        return ""
    if not isinstance(value, (str, int, float, bool)):
        return ""
    text = str(value)
    text = "".join(c for c in text if c in "\n\t" or
                   (ord(c) >= 32 and not 0xD800 <= ord(c) <= 0xDFFF and
                    ord(c) not in {0xFFFE, 0xFFFF}))
    return text if len(text) <= limit else text[:limit - 7] + "…［抜粋］"


def _version(value):
    return _text(value, 200) or MISSING


def _unique(values, limit=MAX_REFS):
    return list(dict.fromkeys(_text(v, 300) for v in values if _text(v, 300)))[:limit]


def _units(text):
    return sum(2 if unicodedata.east_asian_width(c) in {"W", "F"} or c in "MWmw@#%&" else 1 for c in text)


def _wrap(text, width=LINE_UNITS):
    lines = []
    for source_line in _text(text).splitlines() or [""]:
        line, size = "", 0
        for char in source_line:
            count = 2 if unicodedata.east_asian_width(char) in {"W", "F"} or char in "MWmw@#%&" else 1
            if size + count > width:
                lines.append(line)
                line, size = "", 0
            line += char
            size += count
        lines.append(line)
    return lines


def _safe_url(value):
    """Keep verified references as inert metadata. Never fetch them."""
    if not isinstance(value, str) or len(value) > 2048 or any(ord(c) < 33 for c in value):
        return ""
    try:
        parsed = urlsplit(value)
        if parsed.scheme in {"https", "http"} and parsed.netloc and not parsed.username and not parsed.password:
            return value
    except ValueError:
        pass
    return ""


def _saved_source_refs(value):
    refs = []
    for row in _rows(value)[:MAX_REFS]:
        if isinstance(row, str):
            refs.append({"id": _text(row, 300)})
        elif isinstance(row, dict):
            safe = {key: _text(row[key], 500) for key in
                    ("id", "source_id", "note_id", "version", "hash", "source_hash", "label", "title")
                    if isinstance(row.get(key), (str, int, float, bool))}
            url = _safe_url(row.get("url"))
            if url:
                safe["url"] = url
            if safe:
                refs.append(safe)
    return refs


def _evidence_ids(value):
    found, visited = [], 0
    def visit(node, depth=0):
        nonlocal visited
        visited += 1
        if visited > 10000 or depth > 12:
            return
        if isinstance(node, dict):
            for key, child in node.items():
                if key == "evidence_ids":
                    found.extend(v for v in _rows(child) if isinstance(v, str))
                elif key == "evidence_id" and isinstance(child, str):
                    found.append(child)
                elif key in {"claims", "findings", "rows", "previews", "issues", "method"}:
                    visit(child, depth + 1)
        elif isinstance(node, list):
            for child in node[:500]:
                visit(child, depth + 1)
    visit(value)
    return _unique(found)


FIELD_LABELS = {
    "summary": "要約", "summaries": "要約", "text": "本文", "title": "題名", "kind": "区分",
    "claims": "主張", "claim_id": "主張ID", "method": "手法", "method_id": "手法ID", "method_version": "手法版",
    "analysis_unit": "分析単位", "denominator": "分母", "missing_count": "欠測数", "count": "件数",
    "proportion": "割合", "label": "ラベル", "speaker": "話者", "speaker_name": "話者名", "group": "群",
    "term": "語", "surface": "表層形", "lemma": "原形", "pos": "品詞", "frequency": "頻度",
    "segment_count": "発話数", "turns": "交替数", "seconds": "秒", "duration": "長さ", "share": "構成比",
    "ratio": "比率", "percent": "割合（%）", "value": "値", "start": "開始", "end": "終了",
    "rows": "行", "previews": "保存表", "findings": "所見", "limitations": "留保", "issue_id": "指摘ID",
    "issue_key": "指摘キー", "target_id": "対象ID", "target_version": "対象版", "evidence_version": "根拠版",
    "severity": "重大度", "reason": "理由", "missing_evidence": "不足根拠", "alternative": "代替説明",
    "proposed_test": "提案検証（実施状態は履歴参照）", "status": "状態", "stale": "旧版", "review_status": "レビュー状態",
    "reviewed_scope": "レビュー範囲", "response_id": "応答ID", "disposition": "採否", "impact": "影響",
    "decision_id": "判断ID", "iteration": "反復", "view_version": "表示版", "critique_responses": "批判への応答",
    "alternatives": "代替説明", "unresolved": "未解決", "annotation_version": "ラベル版", "codebook_version": "コードブック版",
    "input_hash": "入力hash", "initial_hash": "初期版hash", "initial_id": "初期版ID", "run_id": "Run ID", "item_id": "会話ID",
    "source_revision": "原資料版", "analysis_revision": "分析版", "generation": "実行世代", "last_decision_id": "最終判断ID",
    "result_count": "結果件数", "result_versions_hash": "結果版一覧hash", "result_versions_limit": "結果版一覧の表示上限", "snapshot_signature": "保存版署名",
    "dataset_version": "入力版", "research_mode": "研究モード", "label_field": "ラベル項目", "stop": "停止判断",
}


def _display_path(path):
    path = re.sub(r"[A-Za-z_]+", lambda match: FIELD_LABELS.get(match.group(0), match.group(0)), path)
    path = re.sub(r"\[(\d+)\]", lambda match: str(int(match.group(1)) + 1), path)
    return path.replace(".", " / ")


def _state(value, kind):
    labels = {
        "validation": {"valid": "形式確認済み", "received": "形式検査前", "quarantined": "隔離", "invalid": "形式不適合"},
        "content": {"unreviewed": "内容未確認"},
        "kind": {"observation": "観察", "interpretation": "解釈", "hypothesis": "探索的仮説"},
        "severity": {"high": "高", "medium": "中", "low": "低"},
        "disposition": {"adopt": "採用・解決とは別", "reject": "不採用", "defer": "保留"},
        "status": {"adopted_unresolved": "採用・未解決", "resolved": "解決済み", "completed": "処理完了", "failed": "失敗", "running": "実行中"},
        "review_status": {"issues": "指摘あり", "no_issues": "この範囲では指摘なし", "undetermined": "判定不能", "incomplete": "未完了", "reviewed": "レビュー済み"},
    }
    original = _version(value)
    label = labels.get(kind, {}).get(original)
    return label + "（" + original + "）" if label else original


def _facts(value, prefix="", maximum=40, machine_paths=False):
    """Bounded path-labelled extraction, never a synthetic scientific summary."""
    result = []
    def visit(node, path, depth=0):
        if len(result) >= maximum or depth > 6:
            return
        if isinstance(node, dict):
            # Public analysis contract fields only. Provider traces, prompts,
            # private reasoning and arbitrary legacy fields never enter slides.
            preferred = ("summary", "text", "title", "kind", "claim_id", "method_id", "method_version",
                         "analysis_unit", "denominator", "missing_count", "count", "proportion",
                         "label", "speaker", "speaker_name", "group", "term", "surface", "lemma", "pos",
                         "frequency", "segment_count", "turns", "seconds", "duration", "share", "ratio",
                         "percent", "value", "start", "end", "source", "target", "overlap", "gap",
                         "summaries", "findings", "rows", "previews", "limitations", "claims", "method",
                         "issue_id", "issue_key", "target_id", "target_version", "evidence_version",
                         "severity", "reason", "missing_evidence", "alternative", "proposed_test",
                         "status", "stale", "review_status", "reviewed_scope", "response_id", "disposition",
                         "impact", "decision_id", "iteration", "view_version", "critique_responses",
                         "alternatives", "unresolved", "annotation_version", "codebook_version",
                         "dataset_version", "research_mode", "label_field", "stop")
            keys = [key for key in preferred if key in node]
            for key in keys[:60]:
                visit(node[key], f"{path}.{key}" if path else str(key), depth + 1)
        elif isinstance(node, list):
            if not node:
                return
            if all(isinstance(v, (str, int, float, bool)) or v is None for v in node[:30]):
                result.append(f"{path if machine_paths else _display_path(path)}: " + ", ".join(_text(v, 300) if v is not None else "null" for v in node[:30]))
            else:
                for index, child in enumerate(node[:20]):
                    visit(child, f"{path}[{index}]", depth + 1)
        elif node is not None:
            text = _text(node, 2000)
            if not machine_paths and path.rsplit(".", 1)[-1] in {"kind", "severity", "disposition", "status", "review_status"}:
                text = _state(node, path.rsplit(".", 1)[-1])
            if text:
                result.append(f"{path if machine_paths else _display_path(path)}: {text}" if path else text)
    visit(value, prefix)
    return result[:maximum]


class AnalysisSlidesService:
    def __init__(self, *, export_result: Callable[[str, str], dict],
                 vault_factory: Callable[[], Any], templates: dict | None = None):
        self.export_result = export_result
        self.vault_factory = vault_factory
        self.templates = copy.deepcopy(templates or {})

    def _template(self, template_id):
        # Legacy URL alias selects the sole configured reviewed template.
        if template_id == "research" and template_id not in self.templates and len(self.templates) == 1:
            template_id = next(iter(self.templates))
        if not isinstance(template_id, str) or template_id not in self.templates:
            raise _error("保存済みのスライド設計を選択してください。", "slide_template_unavailable", "template")
        template = copy.deepcopy(self.templates[template_id])
        sections = _mapping(template.get("design")).get("sections")
        if (template.get("id") != template_id or not template.get("version") or
                not isinstance(template.get("prompt"), str) or not template["prompt"].strip() or
                not isinstance(sections, list) or not 1 <= len(sections) <= 8 or
                any(not isinstance(s, dict) or s.get("key") not in SECTION_KEYS or
                    not isinstance(s.get("title"), str) or not s["title"].strip() for s in sections) or
                len({s["key"] for s in sections}) != len(sections)):
            raise _error("スライド設計の登録内容を確認してください。", "slide_template_unavailable")
        if len(canonical(template)) > 64000:
            raise _error("スライド設計が上限を超えています。", "slide_template_unavailable")
        template["hash"] = fingerprint(template)
        return template

    def _snapshot(self, item_id, run_id):
        exported = self.export_result(item_id, run_id)
        if not isinstance(exported, dict):
            raise _error("保存済み分析の形式が正しくありません。", "slide_snapshot_invalid")
        run, initial = _mapping(exported.get("run")), _mapping(exported.get("initial"))
        if run.get("item_id") != item_id or run.get("run_id") != run_id:
            raise LookupError("対象の分析履歴が見つかりません。")
        try:
            encoded = canonical(exported)
        except (ValueError, TypeError, RecursionError):
            raise _error("保存済み分析を安全に読み取れません。", "slide_snapshot_invalid") from None
        if len(encoded) > MAX_EXPORT_BYTES:
            raise _error("保存履歴がスライドの読取上限を超えています。JSONで確認してください。", "slide_size_limit")
        if initial.get("initial_id") and run.get("initial_id") and initial["initial_id"] != run["initial_id"]:
            raise _error("初期版のIDが一致しません。", "slide_snapshot_invalid")
        snapshot = initial.get("snapshot")
        if snapshot is not None and initial.get("hash") and fingerprint(snapshot) != initial["hash"]:
            raise _error("初期版のhashが一致しません。", "slide_snapshot_invalid")
        if isinstance(snapshot, dict) and snapshot.get("input_hash") and run.get("input_hash") != snapshot["input_hash"]:
            raise _error("入力版が一致しません。", "slide_snapshot_invalid")
        for result in _rows(exported.get("raw_results")):
            if not isinstance(result, dict):
                continue
            if result.get("run_id") and result["run_id"] != run_id:
                raise _error("分析結果の所属が一致しません。", "slide_snapshot_invalid")
            if result.get("raw_hash") and fingerprint(result.get("raw")) != result["raw_hash"]:
                raise _error("分析結果のhashが一致しません。", "slide_snapshot_invalid")
        return copy.deepcopy(exported), fingerprint(exported)

    @staticmethod
    def _provenance(exported, signature, result_limit=MAX_RESULTS):
        run, initial = exported["run"], _mapping(exported.get("initial"))
        results = [r for r in _rows(exported.get("raw_results")) if isinstance(r, dict)]
        return {"item_id": run["item_id"], "run_id": run["run_id"], "snapshot_signature": signature,
                "initial_id": _version(initial.get("initial_id")), "initial_hash": _version(initial.get("hash")),
                **{key: _version(run.get(key)) for key in ("input_hash", "source_revision", "analysis_revision",
                    "annotation_version", "codebook_version", "view_version", "generation", "last_decision_id")},
                "result_count": len(results), "result_versions_limit": result_limit,
                "result_versions_hash": fingerprint([{key: result.get(key) for key in
                    ("result_id", "task_id", "raw_hash", "dataset_version", "annotation_version", "validation_status", "content_status")}
                    for result in results]),
                "result_versions": [{key: _version(result.get(key)) for key in
                    ("result_id", "task_id", "raw_hash", "dataset_version", "annotation_version", "validation_status", "content_status")}
                    for result in results[:result_limit]]}

    def catalog(self, item_id, run_id):
        exported, signature = self._snapshot(item_id, run_id)
        templates = []
        for key in self.templates:
            t = self._template(key)
            templates.append({**{k: _text(t.get(k), 500) for k in ("id", "name", "version", "description")},
                              "design": t["design"], "prompt": t["prompt"],
                              "source_references": t.get("source_references", [])})
        design, error = None, ""
        try:
            artifact = self.vault_factory().read_slide_design(item_id=item_id, run_id=run_id)
            if artifact:
                design = {key: artifact.get(key) for key in
                          ("template_id", "template_version", "snapshot_signature", "design_hash")}
                design["matches_snapshot"] = artifact.get("snapshot_signature") == signature
        except (LookupError, AnalysisContractError, StoreConflict):
            error = "保存済みの設計がないか、ノートの変更・削除を確認する必要があります。"
        return {"templates": templates, "design": design, "snapshot_signature": signature,
                "design_error": error, "download": self._availability()}

    def save_design(self, item_id, run_id, template_id="research", expected_snapshot=None):
        template = self._template(template_id)
        exported, signature = self._snapshot(item_id, run_id)
        if expected_snapshot is not None and expected_snapshot != signature:
            raise _error("保存履歴が更新されました。表示を更新してください。", "slide_snapshot_conflict")
        artifact = {"schema_version": SCHEMA_VERSION, "item_id": item_id, "run_id": run_id,
                    "template_id": template["id"], "template_version": template["version"],
                    "design_hash": template["hash"], "snapshot_signature": signature,
                    "provenance": self._provenance(exported, signature, result_limit=80),
                    "design": template["design"], "prompt": template["prompt"],
                    "source_references": template.get("source_references", [])}
        # This explicit POST is the only state-changing boundary in this service.
        receipt = self.vault_factory().publish_slide_design(item_id=item_id, run_id=run_id, artifact=artifact)
        verified = self.vault_factory().read_slide_design(item_id=item_id, run_id=run_id)
        if verified != artifact:
            raise _error("Obsidianの設計保存を確認できません。生成は開始していません。", "slide_design_unverified")
        return {"design": {key: artifact[key] for key in
                           ("template_id", "template_version", "snapshot_signature", "design_hash")},
                "snapshot_signature": signature, "receipt": receipt}

    def _require_design(self, item_id, run_id, template, signature):
        try:
            artifact = self.vault_factory().read_slide_design(item_id=item_id, run_id=run_id)
        except (LookupError, StoreConflict):
            artifact = None
        if not artifact:
            raise _error("先にスライド設計と指示をObsidianへ保存してください。", "slide_design_required")
        if (artifact.get("item_id") != item_id or artifact.get("run_id") != run_id or
                artifact.get("schema_version") != SCHEMA_VERSION or
                artifact.get("template_id") != template["id"] or
                artifact.get("template_version") != template["version"] or
                artifact.get("design_hash") != template["hash"] or
                artifact.get("design") != template["design"] or
                artifact.get("prompt") != template["prompt"] or
                artifact.get("source_references") != template.get("source_references", [])):
            raise _error("保存済みスライド設計が一致しません。設計保存を確認してください。", "slide_design_conflict")
        if artifact.get("snapshot_signature") != signature:
            raise _error("設計保存後に分析履歴が更新されました。設計を保存し直してください。", "slide_snapshot_conflict")
        return artifact

    @staticmethod
    def _availability():
        available = importlib.util.find_spec("pptx") is not None
        return {"available": available, "reason": "" if available else "python-pptxが未導入です。プレビューは利用できます。"}

    def preview(self, item_id, run_id, template_id="research", expected_snapshot=None):
        template = self._template(template_id)
        exported, signature = self._snapshot(item_id, run_id)
        if expected_snapshot is not None and expected_snapshot != signature:
            raise _error("プレビュー後に分析履歴が更新されました。もう一度確認してください。", "slide_snapshot_conflict")
        self._require_design(item_id, run_id, template, signature)
        return self._project(exported, signature, template)

    def presentation(self, item_id, run_id, template_id="research", expected_snapshot=None):
        if not isinstance(expected_snapshot, str) or not re.fullmatch(r"sha256:[a-f0-9]{64}", expected_snapshot):
            raise _error("プレビューで確認した保存版を指定してください。", "slide_snapshot_required", "expected_snapshot")
        preview = self.preview(item_id, run_id, template_id, expected_snapshot)
        if preview.get("review_required"):
            raise _error(preview["download"]["reason"], "slide_review_required")
        if not preview["download"]["available"]:
            raise _error(preview["download"]["reason"], "slide_renderer_unavailable")
        return render_presentation(preview)

    def _project(self, exported, signature, template):
        run = exported["run"]
        initial = _mapping(exported.get("initial"))
        snapshot = _mapping(initial.get("snapshot"))
        view = _mapping(run.get("current_view"))
        config = _mapping(run.get("config"))
        raw_results = [r for r in _rows(exported.get("raw_results")) if isinstance(r, dict)]
        evidence = {str(row.get("evidence_id", row.get("id", ""))): row
                    for row in _rows(snapshot.get("evidence")) if isinstance(row, dict)}
        provenance = self._provenance(exported, signature)
        base_url = "/api/library/" + quote(run["item_id"], safe="") + "/analysis/orchestration/" + quote(run["run_id"], safe="")
        warnings = [DISCLAIMER, "保存済みの内容を決定的に配置した抜粋です。全文は分析履歴を確認してください。"]
        if run.get("status") not in {"completed", "stopped", "cancelled", "failed"}:
            warnings.append("実行途中の固定版です。新しい結果を含めるには設計保存からやり直してください。")
        if not snapshot:
            warnings.append("初期スナップショットは未記録です。保存済みの履歴だけを表示します。")
        if len(raw_results) > MAX_RESULTS:
            warnings.append(f"結果は全{len(raw_results)}件です。版一覧は先頭{MAX_RESULTS}件の抜粋です。")
        slides = []

        def add(section, title, paragraphs, *, result_ids=(), evidence_ids=(), missing=False, max_pages=3):
            paragraphs = [_text(p) for p in paragraphs if _text(p)] or ["未実施／未記録"]
            all_lines = []
            for paragraph in paragraphs:
                all_lines.extend(_wrap(paragraph))
            visual = build_visual(section, title, exported, result_ids, provenance, [s["key"] for s in template["design"]["sections"]])
            if visual:
                def sanitize_visual(node):
                    if isinstance(node, dict):
                        return {key: sanitize_visual(value) for key, value in node.items()}
                    if isinstance(node, list):
                        return [sanitize_visual(value) for value in node]
                    return _text(node) if isinstance(node, str) else node
                visual = sanitize_visual(visual)
                from .analysis_slide_visuals import can_render_visual
                if not can_render_visual(visual):
                    visual = None
            total_pages = 1 if visual else math.ceil(len(all_lines) / MAX_LINES)
            pages = min(max_pages, total_pages)
            result_ids = _unique(result_ids)
            evidence_ids = _unique(evidence_ids)
            source_refs = [{"id": run["run_id"], "label": "保存済み全履歴", "url": base_url + "/export.json"}]
            source_refs += [{"id": rid, "label": "結果 " + rid, "url": base_url + "/results/" + quote(rid, safe="")} for rid in result_ids]
            source_refs += [{"id": eid, "label": "根拠 " + eid, "url": base_url + "/results"} for eid in evidence_ids]
            selected = [r for r in raw_results if r.get("result_id") in result_ids]
            note_provenance = {k: v for k, v in provenance.items() if k != "result_versions"}
            note_provenance["result_versions"] = [{k: r.get(k, MISSING) for k in
                ("result_id", "task_id", "raw_hash", "dataset_version", "annotation_version", "validation_status", "content_status", "stale")}
                for r in selected]
            if section == "evidence":
                note_provenance["result_versions"] = provenance.get("result_versions", [])
            referenced = [{**{key: _version(evidence.get(eid, {}).get(key)) for key in
                ("evidence_id", "utterance_id", "dataset_version", "source_hash", "start", "end", "excluded")},
                "source_refs": _saved_source_refs(evidence.get(eid, {}).get("source_refs"))}
                for eid in evidence_ids]
            for result in selected:
                source_refs.extend(_saved_source_refs(result.get("source_refs")))
                if result.get("validation_status") == "valid":
                    source_refs.extend(_saved_source_refs(_mapping(result.get("raw")).get("source_refs")))
            decision_refs = [{key: _version(d.get(key)) for key in ("decision_id", "result_id", "view_version")}
                             for d in _rows(exported.get("decisions")) if isinstance(d, dict)
                             and d.get("result_id") in result_ids][-12:]
            claim_rows = list(_rows(view.get("claims"))) if section in {"purpose", "major_results", "core"} else []
            for result in selected:
                if result.get("validation_status") == "valid":
                    claim_rows += _rows(_mapping(result.get("raw")).get("claims"))
            claim_refs = [{"claim_id": _version(c.get("claim_id")), "kind": _version(c.get("kind")),
                           "evidence_ids": _unique(_rows(c.get("evidence_ids")))}
                          for c in claim_rows if isinstance(c, dict)][:40]
            note = (DISCLAIMER + "\n" + json.dumps({"provenance": note_provenance, "evidence": referenced,
                "decision_refs": decision_refs, "claim_refs": claim_refs,
                "field_labels": FIELD_LABELS,
                "source_fields": {str(r.get("result_id")): [line.split(": ", 1)[0] for line in _facts(r.get("raw"), maximum=60, machine_paths=True)]
                                  for r in selected if r.get("validation_status") == "valid"},
                "source_refs": source_refs, "saved_fields": paragraphs, "template_id": template["id"], "template_version": template["version"],
                "design_hash": template["hash"]}, ensure_ascii=False, indent=2))
            for index in range(pages):
                lines = paragraphs if visual else all_lines[index * MAX_LINES:(index + 1) * MAX_LINES]
                truncated = total_pages > pages and index == pages - 1
                if truncated:
                    lines[-1] = "［抜粋・全文ではありません。続きは保存履歴へ］"
                slides.append({"id": f"slide-{len(slides) + 1}", "section": section,
                    "title": _text(title, 40) + (f" ({index + 1}/{pages})" if pages > 1 else ""),
                    "paragraphs": lines, "evidence_ids": evidence_ids, "result_ids": result_ids,
                    "source_refs": source_refs, "notes": _text(note, MAX_NOTES),
                    "missing": missing, "truncated": truncated, "visual": visual})

        def record_result(section, title, results, role):
            if not results:
                assigned = [t for t in _rows(run.get("tasks")) if isinstance(t, dict) and t.get("role") == role]
                state = "未実施" if not assigned else "未記録（タスクの状態は履歴を参照）"
                add(section, title, [state], missing=True, max_pages=1)
                return
            lines, eids, rids = [], [], []
            # Most recent saved results fit the template. Earlier results remain linked.
            for result in results[-6:]:
                rid = _text(result.get("result_id")) or MISSING
                rids.append(rid)
                lines.append(f"結果 {rid} / {_state(result.get('validation_status'), 'validation')} / {_state(result.get('content_status'), 'content')}")
                lines.append("この結果のラベル版: " + _version(result.get("annotation_version")) +
                             " / 入力版: " + _version(result.get("dataset_version")))
                if result.get("stale"):
                    lines.append("旧版（stale）の保存結果です")
                if result.get("validation_status") != "valid":
                    lines.append("隔離・未検証のため内容を掲載していません。原記録を参照してください。")
                    continue
                raw = result.get("raw")
                eids += _evidence_ids(raw)
                lines += _facts(raw, maximum=20) or ["未対応形式／本文未記録。元記録を参照してください。"]
            if len(results) > 6:
                lines.append(f"全{len(results)}結果のうち最新6件を抜粋しています。")
            add(section, title, lines, result_ids=rids, evidence_ids=eids, max_pages=6 if role == "critic" else 3)

        for section in template["design"]["sections"]:
            key, title = section["key"], section["title"]
            if key == "purpose":
                question = config.get("question", config.get("objective"))
                add(key, title, [question or "研究目的は未記録", "分析モード: " + _version(config.get("research_mode")),
                    "対象Run: " + run["run_id"], "保存Core要約: " + (_text(view.get("summary")) or "未記録"),
                    "表示版: " + _version(run.get("view_version")) + " / ラベル版: " + _version(run.get("annotation_version")),
                    "保存状態: " + _version(run.get("status")),
                    "終了理由: " + _version(run.get("stop_reason")),
                    "入力変更: " + ("旧版（stale）" if run.get("stale") is True else ("保存記録上の変更なし" if run.get("stale") is False else "未記録")), DISCLAIMER], missing=not bool(question))
            elif key == "method":
                methods = [t for t in _rows(run.get("tasks")) if isinstance(t, dict)]
                lines = ["固定入力: " + _version(run.get("input_hash")),
                         "初期版: " + _version(initial.get("initial_id")),
                         "Handlerの保存済みタスク: " + str(len(methods)),
                         "停止条件: " + _version(config.get("stop_mode")),
                         "終了理由: " + _version(run.get("stop_reason"))]
                if evidence:
                    excluded = sum(row.get("excluded") is True for row in evidence.values())
                    unknown = sum(not isinstance(row.get("excluded"), bool) for row in evidence.values())
                    lines.append(f"保存根拠レコード: {len(evidence)}件 / 除外: {excluded}件 / 採否未記録: {unknown}件")
                else:
                    lines.append("包含・除外の保存根拠レコードは未記録")
                lines += [f"{ROLE_LABELS.get(t.get('role'), _version(t.get('role')))}: {_version(t.get('method_id'))} / {_version(t.get('status'))} / {_text(t.get('title'), 200)}" for t in methods[-12:]]
                add(key, title, lines, max_pages=2)
            elif key == "major_results":
                claims = _rows(view.get("claims"))
                lines = []
                for claim in claims[:24]:
                    if isinstance(claim, dict):
                        lines.append(f"{_state(claim.get('kind'), 'kind')} / {_version(claim.get('claim_id'))}: {_text(claim.get('text'))}")
                        ids = [eid for eid in _rows(claim.get("evidence_ids")) if isinstance(eid, str)]
                        if not ids:
                            lines.append("この保存主張の根拠IDは未記録です")
                        elif any(eid not in evidence or evidence[eid].get("excluded") is not False or
                                 evidence[eid].get("dataset_version") != run.get("input_hash") for eid in ids):
                            lines.append("要確認: 根拠の参照欠落・除外・採否未記録・版不一致があります")
                # Initial data is never recomputed. Unknown legacy structures stay explicitly labelled.
                if not lines:
                    lines = _facts(_mapping(snapshot.get("analysis")).get("research"), "初期保存結果", maximum=16)
                add(key, title, lines or ["主要結果は未記録"], evidence_ids=_evidence_ids(claims), missing=not bool(lines))
            elif key == "specialist_perspectives":
                for role in ("interpretation", "statistics", "verification"):
                    record_result(key, ROLE_LABELS[role], [r for r in raw_results if r.get("role") == role], role)
            elif key == "critic":
                record_result(key, title, [r for r in raw_results if r.get("role") == "critic"], "critic")
                issues = sorted([i for i in _rows(run.get("issues")) if isinstance(i, dict)],
                                key=lambda i: ({"high": 0, "medium": 1, "low": 2}.get(i.get("severity"), 3), i.get("status") == "resolved"))
                if issues:
                    add(key, "批判とCoreの応答", _facts(issues, maximum=40) + _facts(run.get("critique_responses"), "Core応答", maximum=20),
                        evidence_ids=_evidence_ids(issues), result_ids=[i.get("result_id") for i in issues if isinstance(i, dict)], max_pages=4)
            elif key == "core":
                lines = (["Core保存要約: " + _text(view.get("summary"))] if view.get("summary") else ["Core統合は未記録"])
                lines += _facts(view.get("alternatives"), "代替説明", maximum=10)
                lines += _facts(_rows(exported.get("decisions"))[-3:], "保存判断", maximum=24)
                add(key, title, lines, result_ids=[r.get("result_id") for r in raw_results if r.get("role") == "core"][-3:], evidence_ids=_evidence_ids(view), missing=not bool(view))
            elif key == "limitations":
                lines = [DISCLAIMER, "レビュー状態: " + _version(run.get("review_status"))]
                lines += _facts(view.get("unresolved"), "未解決", maximum=12)
                lines += _facts(sorted([i for i in _rows(run.get("unresolved_issues")) if isinstance(i, dict)],
                    key=lambda i: {"high": 0, "medium": 1, "low": 2}.get(i.get("severity"), 3)), "未解決の指摘（重大度順）", maximum=20)
                for result in raw_results[-12:]:
                    if result.get("validation_status") != "valid":
                        continue
                    raw = _mapping(result.get("raw"))
                    lines += _facts(raw.get("limitations", _mapping(raw.get("method")).get("limitations")), "保存された制約", maximum=6)
                    for proposal in _rows(raw.get("analysis_requests"))[:6]:
                        if isinstance(proposal, dict):
                            lines.append("保存された追加分析の依頼案（実施状態はタスク履歴を確認）: " +
                                         _text(proposal.get("question"), 500) + " / role: " + _version(proposal.get("role")) +
                                         " / method: " + _version(proposal.get("method_id")))
                lines.append("未実行の追加検証は提案です。実施結果として扱いません。")
                add(key, title, lines, max_pages=5)
            elif key == "evidence":
                lines = [f"{FIELD_LABELS.get(key, key)}: {value}" for key, value in provenance.items() if key != "result_versions"]
                referenced_ids = _unique([eid for slide in slides for eid in slide["evidence_ids"]])
                for eid in referenced_ids[:20]:
                    ev = evidence.get(eid)
                    if not ev:
                        lines.append(eid + ": 根拠本文は未記録")
                    elif ev.get("excluded"):
                        lines.append(eid + ": 除外済み（本文非掲載）")
                    else:
                        lines.append(f"{eid} / 発話 {_version(ev.get('utterance_id'))}: {_text(ev.get('text'), 240)}")
                add(key, title, lines, evidence_ids=referenced_ids, max_pages=3)
        if len(slides) > MAX_SLIDES:
            raise _error("スライド構成がページ上限を超えました。設計を確認してください。", "slide_size_limit")
        if any(s["truncated"] for s in slides):
            warnings.append("長い節はページ上限で抜粋されています。各ページの根拠リンクから全文を確認できます。")
        review_required = any(s["truncated"] and s["section"] in {"critic", "limitations"} for s in slides)
        critical_body = "".join(line for slide in slides if slide["section"] in {"critic", "limitations"}
                                for line in ([visible_text(slide["visual"])] if slide.get("visual") else slide["paragraphs"]))
        for issue in _rows(run.get("issues")):
            if not isinstance(issue, dict) or issue.get("severity") != "high" or issue.get("status") == "resolved":
                continue
            # A source link cannot substitute for an omitted material caveat.
            # Compare saved text exactly (line wrapping excluded), never judge it
            # with a model or treat adoption/staleness as resolution.
            required = [issue.get("issue_id"), issue.get("reason")]
            required += [issue.get(key) for key in ("missing_evidence", "alternative", "proposed_test") if issue.get(key)]
            if any(not isinstance(value, str) or not value or len(value) > 4000 or
                   "".join(value.splitlines()) not in critical_body for value in required):
                review_required = True
        reference_issues = []
        used_evidence = _unique([eid for slide in slides for eid in slide["evidence_ids"]])
        for eid in used_evidence:
            row = evidence.get(eid)
            if not row:
                reference_issues.append(eid + ": 根拠の参照欠落")
            elif (row.get("excluded") is not False or row.get("dataset_version") != run.get("input_hash") or
                  not isinstance(row.get("text"), str) or row.get("source_hash") != fingerprint(row.get("text"))):
                reference_issues.append(eid + ": 除外・採否未記録・入力版・原文hashを要確認")
        used_results = {rid for slide in slides for rid in slide["result_ids"]}
        for result in raw_results:
            if result.get("result_id") in used_results and result.get("validation_status") == "valid" and (
                    not result.get("dataset_version") or result.get("dataset_version") != run.get("input_hash")):
                reference_issues.append(_version(result.get("result_id")) + ": 結果の入力版が不明または不一致")
        if reference_issues:
            review_required = True
            warnings.extend(reference_issues[:20])
        download = self._availability()
        if review_required:
            reason = ("根拠の参照・入力版・原文hashを確認できません。保存履歴を確認してからPPTXを作成してください。"
                      if reference_issues else "批判・限界の抜粋、または重大な未解決指摘の掲載不足があります。重要な留保の脱落を避けるため、設計を見直してからPPTXを作成してください。")
            warnings.append(reason)
            download = {"available": False, "reason": reason}
        return {"schema_version": SCHEMA_VERSION, "template": {key: template.get(key, "") for key in ("id", "name", "version")},
                "snapshot_signature": signature, "provenance": provenance, "slides": slides,
                "warnings": warnings, "synthetic": config.get("synthetic") is True, "limits": {"max_slides": MAX_SLIDES, "max_lines": MAX_LINES, "line_units": LINE_UNITS},
                "download": download, "review_required": review_required, "style": _mapping(template["design"].get("style"))}


def render_presentation(preview):
    """Render native editable diagrams/charts/tables; no remote assets or input code."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except ImportError:
        raise _error("python-pptxが未導入です。プレビューは利用できます。", "slide_renderer_unavailable") from None
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = "Gurumoji 保存済み分析"
    prs.core_properties.subject = preview["provenance"]["run_id"]
    prs.core_properties.author = "Gurumoji"
    prs.core_properties.keywords = preview["snapshot_signature"]
    # Fixed metadata avoids introducing a misleading analysis/generation timestamp.
    prs.core_properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    prs.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    style = _mapping(preview.get("style"))
    font = _text(style.get("font"), 100) or "Yu Gothic"

    def color(key, fallback):
        value = _text(style.get(key), 20).lstrip("#")
        return value.upper() if re.fullmatch(r"[a-fA-F0-9]{6}", value) else fallback

    def box(slide, x, y, w, h, lines, size, color):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0)
        frame.margin_top = frame.margin_bottom = Inches(0)
        for index, line in enumerate(lines):
            p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            p.text = _text(line)
            p.font.name, p.font.size = font, Pt(size)
            from pptx.oxml.xmlchemy import OxmlElement
            east_asian = OxmlElement("a:ea")
            east_asian.set("typeface", font)
            p.font._element.append(east_asian)
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_before, p.space_after = Pt(0), Pt(3)
            p.line_spacing = 1.12
        return shape

    for index, dto in enumerate(preview["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(color("background_color", "FFFEF9"))
        title_lines = _wrap(dto["title"], 52)[:2]
        box(slide, .65, .4, 12, 1.1, title_lines, 32, color("title_color", "1C6B50"))
        from .analysis_slide_visuals import render_visual_slide
        rendered = render_visual_slide(prs, slide, dto, font=font, colors={
            "title": color("title_color", "1C6B50"), "body": color("body_color", "18211D"),
            "muted": color("footer_color", "68736E"), "accent": "A44B37", "background": color("background_color", "F4F5F0")})
        if dto.get("visual") and not rendered:
            raise _error("視覚レイアウトを安全に配置できません。設計を確認してください。", "slide_layout_overflow")
        if not rendered:
            box(slide, .65, 1.65, 12, 4.9, dto["paragraphs"], 22, color("body_color", "18211D"))
        footer = (f"{index + 1}/{len(preview['slides'])}  " + ("匿名の合成サンプル  " if preview.get("synthetic") else "探索分析・抜粋  ") +
                  f"run:{_text(preview['provenance']['run_id'], 20)}  "
                  f"view:{preview['provenance']['view_version']} / label:{preview['provenance']['annotation_version']}  "
                  f"{preview['snapshot_signature'][7:19]}")
        box(slide, .65, 6.83, 12, .3, [footer], 14, color("footer_color", "586875"))
        slide.notes_slide.notes_text_frame.text = dto["notes"]
    stream = io.BytesIO()
    prs.save(stream)
    stream.seek(0)
    # OOXML is ZIP: normalize package timestamps/order so identical saved data
    # and template yield identical bytes in the same installed renderer version.
    stable = io.BytesIO()
    with ZipFile(stream) as source, ZipFile(stable, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            payload = source.read(name)
            if name.startswith("ppt/embeddings/") and name.endswith(".xlsx"):
                # Native chart workbooks otherwise contain generation-time
                # metadata even when the saved run and chart values are fixed.
                workbook = io.BytesIO()
                with ZipFile(io.BytesIO(payload)) as original, ZipFile(workbook, "w", compression=ZIP_DEFLATED) as normalized:
                    for member in sorted(original.namelist()):
                        data = original.read(member)
                        if member == "docProps/core.xml":
                            data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+",
                                          rb"\g<1>2000-01-01T00:00:00Z", data)
                        entry = ZipInfo(member, date_time=(2000, 1, 1, 0, 0, 0))
                        entry.compress_type = ZIP_DEFLATED
                        normalized.writestr(entry, data)
                payload = workbook.getvalue()
            target.writestr(info, payload)
    stable.seek(0)
    return stable
