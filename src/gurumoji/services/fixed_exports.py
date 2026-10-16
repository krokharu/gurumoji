"""Exports from verified, typed saved tables; never reruns an analysis."""
from __future__ import annotations

import hashlib
import html
import io
import json
from pathlib import Path
import tempfile
import zipfile

from .group_analysis import analysis_csv_safe


def tables_for_run(store, run_id):
    snapshot, result, manifest, _ = store.readable_package(run_id)
    tables = {name: store.read_table(run_id, name) for name in manifest.get('typed_tables', {})}
    if not tables:
        raise ValueError('型付き全件表のない旧保存結果には対応していません。')
    provenance = {'run_id': run_id, 'snapshot_id': store.get(run_id)['snapshot_id'],
        'input_hash': hashlib.sha256(json.dumps(snapshot, ensure_ascii=False, sort_keys=True,
            separators=(',', ':'), allow_nan=False).encode()).hexdigest(),
        'methods': result.get('methods', []), 'algorithms': result.get('algorithms', {}),
        'source_revision': store.get(run_id)['source_revision'],
        'analysis_revision': store.get(run_id)['analysis_revision']}
    return tables, provenance


def text_cell(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else value


def workbook_bytes(tables, provenance):
    from openpyxl import Workbook
    book = Workbook()
    meta = book.active
    meta.title = 'README'
    meta.append(['固定保存結果', provenance['run_id']])
    provenance_text = json.dumps(provenance, ensure_ascii=False)
    for position in range(0, len(provenance_text), 30000):
        meta.append(['来歴' if position == 0 else '来歴（続き）', analysis_csv_safe(provenance_text[position:position+30000])])
    meta.append(['欠測', '_missingシートでnull・空文字・未存在を区別します。'])
    missing = book.create_sheet('_missing')
    missing.append(['dataset', 'row_id', 'column', 'state'])
    used = {'README', '_missing'}
    for name, table in tables.items():
        title = name[:31]
        if title in used:
            title = name[:22] + '-' + hashlib.sha256(name.encode()).hexdigest()[:8]
        used.add(title)
        sheet = book.create_sheet(title)
        if len(table['rows']) > 1048575 or len(table['fields']) > 16384:
            raise ValueError('Excelの行数・列数上限を超えています。JSONを使用してください。')
        sheet.append([analysis_csv_safe(field) for field in table['fields']])
        sheet.freeze_panes = 'A2'
        for index, row in enumerate(table['rows']):
            values = row['values']
            cells = []
            for field in table['fields']:
                value = values.get(field)
                if field not in values or value is None or value == '':
                    state = 'absent' if field not in values else 'null' if value is None else 'empty_string'
                    missing.append([analysis_csv_safe(value) for value in (name, row.get('row_id', str(index)), field, state)])
                value = text_cell(value)
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError('Excelのセル文字数上限を超えています。JSONを使用してください。')
                cells.append(analysis_csv_safe(value) if isinstance(value, str) else value)
            sheet.append(cells)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def sav_bundle(tables, provenance):
    import pandas as pd
    import pyreadstat
    stream = io.BytesIO()
    manifest = {**provenance, 'format': 'sav-bundle-1', 'datasets': {},
        'note': 'SAVの型制約を補うため、元の型付きJSONを同梱します。欠測と空文字はJSONを参照してください。'}
    with tempfile.TemporaryDirectory(prefix='gurumoji-sav-') as temporary, zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, table in tables.items():
            columns, labels = {}, {}
            for index, field in enumerate(table['fields']):
                variable = f'v{index + 1:05d}'
                labels[variable] = field
                values = [row['values'].get(field) for row in table['rows']]
                numeric = all(value is None or type(value) in (int, float, bool) for value in values)
                columns[variable] = ([float(value) if value is not None else float('nan') for value in values]
                    if numeric else ['' if value is None else str(text_cell(value)) for value in values])
                if not numeric and any(len(v.encode('utf-8')) > 32767 for v in columns[variable]):
                    raise ValueError('SAVの文字列上限を超えています。JSONを使用してください。')
            if not columns:
                continue
            path = Path(temporary) / (name + '.sav')
            pyreadstat.write_sav(pd.DataFrame(columns), str(path), column_labels=labels,
                file_label='Gurumoji fixed run ' + provenance['run_id'])
            data = path.read_bytes()
            archive.writestr(name + '.sav', data)
            archive.writestr(name + '.typed.json', json.dumps(table, ensure_ascii=False, allow_nan=False))
            manifest['datasets'][name] = {'columns': labels, 'rows': len(table['rows']),
                'sha256': hashlib.sha256(data).hexdigest()}
        archive.writestr('manifest.json', json.dumps(manifest, ensure_ascii=False, allow_nan=False))
    return stream.getvalue()


def chart_values(table):
    fields = table['fields']
    numeric = next((field for field in fields if any(type(r['values'].get(field)) in (int, float)
        for r in table['rows'])), None)
    if numeric is None:
        raise ValueError('この表には数値列がありません。')
    label = next((field for field in fields if field != numeric and any(isinstance(r['values'].get(field), str)
        for r in table['rows'])), None)
    return numeric, [(str(row['values'].get(label, index + 1)), row['values'].get(numeric))
                     for index, row in enumerate(table['rows'][:30])]


def image_bytes(table, name, provenance, format):
    numeric, values = chart_values(table)
    width, height = 1000, max(220, 120 + len(values) * 30)
    maximum = max([abs(v) for _, v in values if type(v) in (int, float)] or [1]) or 1
    caption = f"{name} / {numeric} / 先頭{len(values)}件・全{len(table['rows'])}件 / run {provenance['run_id']}"
    if format == 'svg':
        lines = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" role="img">',
            '<title>' + html.escape(caption) + '</title><rect width="100%" height="100%" fill="#fffdf7"/>',
            '<text x="20" y="30" font-size="16">' + html.escape(caption) + '</text>']
        for index, (label, value) in enumerate(values):
            y = 65 + index * 30
            lines.append(f'<text x="20" y="{y}" font-size="14">{html.escape(label[:30])}</text>')
            if type(value) in (int, float):
                size = abs(value) / maximum * 300
                x = 620 - size if value < 0 else 620
                lines.append(f'<rect x="{x}" y="{y-16}" width="{size}" height="20" fill="#1e624e"/>')
            lines.append(f'<text x="925" y="{y}" font-size="13">{html.escape("欠測" if value is None else str(value))}</text>')
        return ''.join(lines + ['</svg>']).encode('utf-8')
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new('RGB', (width, height), '#fffdf7')
    draw = ImageDraw.Draw(image)
    font_path = Path('C:/Windows/Fonts/meiryo.ttc')
    font = ImageFont.truetype(str(font_path), 15) if font_path.is_file() else ImageFont.load_default(size=15)
    draw.text((20, 20), caption, fill='#163c32', font=font)
    for index, (label, value) in enumerate(values):
        y = 60 + index * 30
        draw.text((20, y), label[:30], fill='#163c32', font=font)
        if type(value) in (int, float):
            size = abs(value) / maximum * 300
            draw.rectangle((620 - size if value < 0 else 620, y, 620 if value < 0 else 620 + size, y + 20), fill='#1e624e')
        draw.text((925, y), '欠測' if value is None else str(value), fill='#163c32', font=font)
    stream = io.BytesIO()
    image.save(stream, format='PNG')
    return stream.getvalue()


def export_fixed(store, run_id, format, dataset=None):
    tables, provenance = tables_for_run(store, run_id)
    if format == 'xlsx':
        return workbook_bytes(tables, provenance), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'
    if format == 'sav':
        return sav_bundle(tables, provenance), 'application/zip', 'sav.zip'
    if format in {'svg', 'png'}:
        if dataset not in tables:
            raise ValueError('画像にする表をdatasetで指定してください。')
        return image_bytes(tables[dataset], dataset, provenance, format), 'image/' + ('svg+xml' if format == 'svg' else 'png'), format
    raise ValueError('未対応の出力形式です。')
