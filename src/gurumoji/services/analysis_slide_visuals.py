"""Native visual layouts for the saved-analysis slide renderer.

This is a presentation adapter, not an analysis step. All text and chart values
come from the already-reviewed projection. It neither reads files nor calls a
model or network service. Unsupported or overfull layouts return ``False``
*before* touching the slide, allowing the caller's complete text fallback.
"""
from __future__ import annotations

import math
import re
import unicodedata


_BOUNDS = (.65, 1.55, 12.0, 4.9)
_DEFAULTS = {"background": "F4F5F0", "title": "1C6B50", "body": "18211D",
             "muted": "68736E", "accent": "A44B37", "rule": "B8C3BA"}


class _DoesNotFit(ValueError):
    """The lossless text fallback is preferable to this visual layout."""


def _plain(value):
    if value is None:
        return ""
    if not isinstance(value, (str, int, float, bool)):
        raise _DoesNotFit("unsupported text")
    result = str(value)
    if len(result) > 8000:
        raise _DoesNotFit("overlong text")
    return "".join(c for c in result if c in "\n\t" or
                   (ord(c) >= 32 and not 0xD800 <= ord(c) <= 0xDFFF and
                    ord(c) not in {0xFFFE, 0xFFFF}))


def _list(value):
    if value is None:
        return []
    if not isinstance(value, list) or len(value) > 50:
        raise _DoesNotFit("unsupported list")
    return [_plain(item) for item in value if _plain(item)]


def _advance(char):
    # Conservative Japanese/Latin advance estimates, in ems. Explicit wrapping
    # avoids relying on PowerPoint's platform-specific automatic font fitting.
    if unicodedata.combining(char):
        return 0
    if unicodedata.east_asian_width(char) in {"W", "F"}:
        return 1.05
    if char in "MWmw@#%&":
        return .85
    return .60


def _wrapped(value, width, height, size):
    text = _plain(value)
    if not text:
        return ""
    available = (width * 72 - 4) / size
    lines = []
    for source in text.replace("\t", "  ").split("\n"):
        line, used = "", 0.0
        for char in source:
            advance = _advance(char)
            if line and used + advance > available:
                lines.append(line)
                line, used = "", 0.0
            line += char
            used += advance
        lines.append(line)
    if len(lines) * size * 1.24 + 3 > height * 72:
        raise _DoesNotFit("text does not fit at its readable font size")
    return "\n".join(lines)


def _palette(colors):
    colors = colors if isinstance(colors, dict) else {}
    aliases = {"muted": "footer_color"}
    palette = {}
    for key, default in _DEFAULTS.items():
        value = colors.get(key, colors.get(aliases.get(key, key + "_color"), default))
        value = str(value).lstrip("#")
        palette[key] = value.upper() if re.fullmatch(r"[a-fA-F0-9]{6}", value) else default
    return palette


class _Scene:
    """Preflight a complete scene before creating any native slide objects."""

    def __init__(self):
        self.items = []

    @staticmethod
    def bounds(x, y, w, h):
        left, top, width, height = _BOUNDS
        if (w < 0 or h < 0 or x < left - .001 or y < top - .001 or
                x + w > left + width + .001 or y + h > top + height + .001):
            raise _DoesNotFit("object outside content area")

    def text(self, value, x, y, w, h, size=22, color="body", bold=False):
        self.bounds(x, y, w, h)
        text = _wrapped(value, w, h, size)
        if text:
            self.items.append(("text", (text, x, y, w, h, size, color, bold)))

    def rule(self, x1, y1, x2, y2, color="rule", width=1, arrow=False):
        self.bounds(min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
        self.items.append(("rule", (x1, y1, x2, y2, color, width, arrow)))

    def metadata(self, values, y=6.02, h=.4):
        self.text("   ".join(_list(values)), .65, y, 12, h, 14, "muted")

    def paragraphs(self, values, x, y, w, h, size=20, color="body"):
        self.text("\n".join(_list(values)), x, y, w, h, size, color)


def _hero(scene, visual):
    scene.text("研究の問い", .65, 1.58, 12, .35, 16, "muted")
    scene.text(visual.get("question"), .65, 2.03, 11.7, 1.06, 26, "title", True)
    scene.rule(.65, 3.34, 12.65, 3.34)
    scene.text("保存された要旨", .65, 3.62, 12, .35, 16, "muted")
    scene.text(visual.get("summary"), .65, 4.07, 11.7, 1.3, 24)
    scene.text(visual.get("caveat"), .65, 5.55, 12, .45, 16, "accent")
    scene.metadata(visual.get("states"), 6.08, .34)


def _method(scene, visual):
    nodes = visual.get("nodes")
    if not isinstance(nodes, list) or not 2 <= len(nodes) <= 4:
        raise _DoesNotFit("method needs two to four saved stages")
    gap = .48
    width = (12 - gap * (len(nodes) - 1)) / len(nodes)
    for index, node in enumerate(nodes):
        if not isinstance(node, dict):
            raise _DoesNotFit("invalid stage")
        x = .65 + index * (width + gap)
        scene.text(str(index + 1).zfill(2), x, 1.92, width, .62, 32, "title", True)
        scene.text(node.get("title"), x, 2.74, width, .8, 20, "title", True)
        scene.rule(x, 3.68, x + width, 3.68, "title", 1.7)
        scene.text(node.get("text"), x, 3.94, width, 1.55, 20)
        if index < len(nodes) - 1:
            scene.rule(x + width + .08, 3.68, x + width + gap - .08, 3.68,
                       "muted", 1.5, True)
    scene.text(visual.get("note"), .65, 5.62, 12, .81, 14, "muted")


def _evidence_chain(scene, visual):
    scene.text("保存された根拠", .65, 1.7, 4.75, .35, 16, "muted")
    scene.rule(.68, 2.23, .68, 4.82, "title", 2.5)
    scene.text(visual.get("quote"), .94, 2.23, 4.28, 2.58, 22)
    scene.text(visual.get("evidence_label"), .94, 5.0, 4.55, .63, 14, "muted")
    scene.text("参照", 5.58, 2.81, .76, .38, 16, "muted")
    scene.rule(5.55, 3.35, 6.36, 3.35, "title", 1.7, True)
    scene.text(visual.get("kind_label") or "保存された主張", 6.7, 1.7, 5.95, .55,
               18, "title", True)
    scene.text(visual.get("claim"), 6.7, 2.42, 5.92, 2.45, 24)
    scene.text(visual.get("claim_id"), 6.7, 5.0, 5.95, .63, 14, "muted")
    scene.text(visual.get("caveat"), .65, 5.84, 12, .58, 18, "accent")


def _perspective(scene, visual):
    scene.text(visual.get("summary"), .65, 1.75, 11.85, 1.02, 24, "title", True)
    scene.rule(.65, 3.03, 12.65, 3.03)
    scene.text(visual.get("kind_label") or "保存された主張", .65, 3.27, 7.65, .56,
               18, "title", True)
    scene.text(visual.get("claim"), .65, 3.99, 7.65, 1.29, 22)
    scene.text("根拠参照", 9.08, 3.27, 3.57, .4, 16, "muted")
    scene.text(visual.get("evidence_label"), 9.08, 3.92, 3.57, 1.35, 16, "muted")
    scene.text(visual.get("caveat"), .65, 5.33, 12, .47, 16, "accent")
    meta = _list(visual.get("meta"))
    if len(meta) > 3:
        scene.text("   ".join(meta[:3]) + "\n" + "   ".join(meta[3:]), .65, 5.88, 12, .55, 14, "muted")
    else:
        scene.metadata(meta, 5.88, .55)


def _statistics(scene, visual):
    categories, values = visual.get("categories"), visual.get("values")
    if (not isinstance(categories, list) or not isinstance(values, list) or
            not 1 <= len(categories) <= 8 or len(categories) != len(values)):
        raise _DoesNotFit("unsupported chart dimensions")
    categories = [_plain(value) for value in categories]
    if any(not value or "\n" in value or sum(map(_advance, value)) > 15 for value in categories):
        raise _DoesNotFit("chart labels need a complete table instead")
    if any(isinstance(value, bool) or not isinstance(value, (int, float)) or
           not math.isfinite(value) or value < 0 for value in values):
        raise _DoesNotFit("chart values must be saved nonnegative finite numbers")
    counts = visual.get("counts")
    if counts is not None:
        if (not isinstance(counts, list) or len(counts) != len(categories) or
                any(isinstance(value, bool) or not isinstance(value, (int, float)) or
                    not math.isfinite(value) or value < 0 for value in counts)):
            raise _DoesNotFit("counts must be saved finite nonnegative values")
        denominator = _plain(visual.get("denominator")) or "未記録"
        categories = [label + "（" + _plain(count) + " / " + denominator + "）"
                      for label, count in zip(categories, counts)]
        if any(sum(map(_advance, label)) > 20 for label in categories):
            raise _DoesNotFit("count-labelled category needs a table")
    unit = _plain(visual.get("unit")) or "件数"
    scene.text(visual.get("summary"), .65, 1.62, 12, .85, 22)
    scene.text(unit, .65, 2.55, 8.05, .37, 16, "muted")
    chart_height = min(2.59, .43 * len(values) + .6)
    chart_top = 2.95 + (2.59 - chart_height) / 2
    scene.items.append(("chart", (categories, list(values), unit, .65, chart_top, 8.05, chart_height)))
    scene.text("分母", 9.3, 2.56, 3.35, .38, 16, "muted")
    scene.text(visual.get("denominator"), 9.3, 3.03, 3.35, .6, 28, "title", True)
    scene.text("欠測", 9.3, 3.86, 3.35, .38, 16, "muted")
    scene.text(visual.get("missing_count"), 9.3, 4.33, 3.35, .6, 28, "accent", True)
    scene.text("ラベル版 " + (_plain(visual.get("annotation_version")) or "未記録"),
               9.3, 5.12, 3.35, .4, 16, "muted")
    scene.paragraphs(visual.get("caveats"), .65, 5.55, 12, .35, 16, "accent")
    scene.metadata(visual.get("meta"), 5.9, .53)


def _critique(scene, visual):
    issue = visual.get("issue") or {}
    response = visual.get("response") or {}
    if not isinstance(issue, dict) or not isinstance(response, dict):
        raise _DoesNotFit("invalid critique")
    scene.text("批判者の指摘", .65, 1.65, 5.5, .45, 18, "accent", True)
    scene.text(issue.get("reason"), .65, 2.29, 5.35, 1.02, 22)
    scene.text("不足している根拠", .65, 3.57, 5.35, .38, 16, "muted")
    scene.text(issue.get("missing_evidence"), .65, 4.01, 5.35, .75, 20)
    scene.text("提案された検証", .65, 4.84, 5.35, .38, 16, "muted")
    scene.text(issue.get("proposed_test"), .65, 5.28, 5.35, .72, 18)
    scene.rule(6.22, 2.75, 6.98, 2.75, "muted", 1.7, True)
    scene.text("Coreの保存済み応答", 7.25, 1.65, 5.4, .45, 18, "title", True)
    scene.text(response.get("disposition") or "応答未記録", 7.25, 2.29, 5.4, .8,
               24, "title", True)
    scene.text(response.get("reason"), 7.25, 3.26, 5.4, .88, 22)
    scene.text("判断への影響", 7.25, 4.42, 5.4, .38, 16, "muted")
    scene.text(response.get("impact"), 7.25, 4.86, 5.4, .73, 20)
    scene.text(visual.get("status"), 7.25, 5.7, 5.4, .38, 18, "accent", True)
    scene.metadata(issue.get("metadata"), 6.13, .3)


def _core(scene, visual):
    scene.text(visual.get("summary"), .65, 1.78, 11.8, 1.09, 24, "title", True)
    scene.rule(.65, 3.12, 12.65, 3.12)
    scene.text("保存された代替説明", .65, 3.43, 5.5, .48, 18, "accent", True)
    scene.paragraphs([text.replace("。", "。\n").rstrip("\n") for text in _list(visual.get("alternatives"))], .65, 4.12, 5.5, 2.16, 20)
    scene.text("批判への応答", 7.05, 3.43, 5.6, .48, 18, "title", True)
    scene.paragraphs(visual.get("responses"), 7.05, 4.12, 5.6, 2.16, 20)


def _limitations(scene, visual):
    scene.text("未解決点", .65, 1.8, 5.5, .46, 20, "accent", True)
    scene.paragraphs(visual.get("unresolved"), .65, 2.56, 5.5, 2.79, 22)
    scene.rule(6.5, 1.85, 6.5, 5.35)
    scene.text("解釈の留保", 7.08, 1.8, 5.57, .46, 20, "title", True)
    scene.paragraphs(visual.get("caveats"), 7.08, 2.56, 5.57, 2.79, 19)
    scene.text(visual.get("note"), .65, 5.72, 12, .71, 18, "muted")


def _audit(scene, visual):
    rows = visual.get("rows")
    if not isinstance(rows, list) or not 1 <= len(rows) <= 10:
        raise _DoesNotFit("audit needs one to ten key/value rows")
    prepared, heights = [["項目", "保存値"]], [.4]
    for row in rows:
        if not isinstance(row, list) or len(row) != 2:
            raise _DoesNotFit("audit rows must contain a key and value")
        row = [_plain(value) for value in row]
        height = .40
        while height <= 1.3:
            try:
                cells = [_wrapped(row[0], 3.85, height - .09, 15),
                         _wrapped(row[1], 7.8, height - .09, 15)]
                break
            except _DoesNotFit:
                height += .25
        else:
            raise _DoesNotFit("audit cell needs a continuation page")
        prepared.append(cells)
        heights.append(height)
    if sum(heights) > 4.42:
        raise _DoesNotFit("audit table needs a continuation page")
    scene.items.append(("table", (prepared, heights, .65, 1.63, 12)))
    scene.text(visual.get("note"), .65, 6.12, 12, .31, 14, "muted")


_LAYOUTS = {"hero": _hero, "method": _method, "evidence_chain": _evidence_chain,
            "perspective": _perspective, "statistics": _statistics,
            "critique": _critique, "core": _core, "limitations": _limitations,
            "audit": _audit}


def _set_font(font_object, name, size, color, bold=False):
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Pt

    font_object.name, font_object.size = name, Pt(size)
    font_object.bold = bold
    font_object.color.rgb = RGBColor.from_string(color)
    # python-pptx exposes only the Latin face. Explicit East Asian typeface
    # prevents Office from selecting a different Japanese face for native text.
    properties = font_object._rPr
    east_asian = properties.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        properties.append(east_asian)
    east_asian.set("typeface", name)


def _paint_text(slide, item, font, colors):
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    text, x, y, w, h, size, color, bold = item
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        _set_font(paragraph.font, font, size, colors[color], bold)
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.18
    return shape


def _paint_rule(slide, item, colors):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    x1, y1, x2, y2, color, width, arrow = item
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                       Inches(x2), Inches(y2))
    shape.line.color.rgb = RGBColor.from_string(colors[color])
    shape.line.width = Pt(width)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        end.set("w", "sm")
        end.set("len", "sm")
        shape.line._get_or_add_ln().append(end)


def _paint_chart(slide, item, font, colors):
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK
    from pptx.util import Inches, Pt

    categories, values, unit, x, y, w, h = item
    data = CategoryChartData()
    data.categories = categories
    data.add_series(unit, values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y),
                                   Inches(w), Inches(h), data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name, chart.font.size = font, Pt(16)
    chart.font.color.rgb = RGBColor.from_string(colors["muted"])
    plot = chart.plots[0]
    plot.gap_width = 58 if len(values) <= 4 else 42
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    labels.show_value = True
    labels.show_category_name = labels.show_series_name = labels.show_legend_key = False
    labels.number_format = "0" if all(float(v).is_integer() for v in values) else "0.###"
    _set_font(labels.font, font, 18, colors["body"])
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor.from_string(colors["title"])
    series.format.line.fill.background()
    category_axis, value_axis = chart.category_axis, chart.value_axis
    category_axis.reverse_order = True
    _set_font(category_axis.tick_labels.font, font, 18, colors["body"])
    _set_font(value_axis.tick_labels.font, font, 16, colors["muted"])
    value_axis.minimum_scale = 0
    # A small right margin fits native value labels without changing the values.
    value_axis.maximum_scale = 100 if unit == "割合（%）" else max(values) * 1.22 if max(values) > 0 else 1
    if unit == "割合（%）":
        value_axis.major_unit = 25
    value_axis.has_major_gridlines = False
    value_axis.tick_labels.number_format = labels.number_format
    for axis in (category_axis, value_axis):
        axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.color.rgb = RGBColor.from_string(colors["rule"])
        axis.format.line.width = Pt(.75)
    # Axis IDs are OOXML unsignedInt. python-pptx's built-in chart template uses
    # signed IDs, which Office tolerates but strict consumers cannot parse.
    # Normalize every reference with the same bijection, retaining the links.
    for element in chart._chartSpace.iter():
        if element.tag.rsplit("}", 1)[-1] in {"axId", "crossAx"}:
            element.set("val", str(int(element.get("val")) % (2 ** 32)))


def _paint_table(slide, item, font, colors):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    rows, heights, x, y, width = item
    table = slide.shapes.add_table(len(rows), 2, Inches(x), Inches(y),
                                   Inches(width), Inches(sum(heights))).table
    table.columns[0].width, table.columns[1].width = Inches(4.05), Inches(7.95)
    table.first_row = False
    table.horz_banding = False
    for row_index, (row, height) in enumerate(zip(rows, heights)):
        table.rows[row_index].height = Inches(height)
        for column_index, text in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(colors["background"])
            cell.margin_left, cell.margin_right = Inches(.08), Inches(.07)
            cell.margin_top = cell.margin_bottom = Inches(.045)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = text
            for paragraph in cell.text_frame.paragraphs:
                _set_font(paragraph.font, font, 15,
                          colors["title"] if row_index == 0 or column_index == 0 else colors["body"],
                          row_index == 0)
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.12
            properties = cell._tc.get_or_add_tcPr()
            # A flat typographic table: horizontal rules, no panel fills or grid.
            for edge in ("L", "R", "T", "B"):
                line = OxmlElement("a:ln" + edge)
                line.set("w", "12700" if row_index == 0 else "6350")
                if edge == "B":
                    fill = OxmlElement("a:solidFill")
                    color = OxmlElement("a:srgbClr")
                    color.set("val", colors["title"] if row_index == 0 else colors["rule"])
                    fill.append(color)
                    line.append(fill)
                else:
                    line.append(OxmlElement("a:noFill"))
                properties.append(line)


def can_render_visual(visual) -> bool:
    """Run the exact lossless-layout preflight without creating a PPTX object."""
    if not isinstance(visual, dict) or not isinstance(visual.get("kind"), str):
        return False
    layout = _LAYOUTS.get(visual["kind"])
    if layout is None:
        return False
    try:
        layout(_Scene(), visual)
    except _DoesNotFit:
        return False
    return True


def render_visual_slide(prs, slide, dto, *, font, colors) -> bool:
    """Draw a supported visual DTO without cropping or inventing its content.

    ``False`` means no native objects were added. The caller must retain its
    existing text fallback, title, footer and provenance notes. The ``prs``
    argument keeps the renderer interface explicit; it is never saved here.
    """
    del prs
    visual = dto.get("visual") if isinstance(dto, dict) else None
    if not isinstance(visual, dict) or not isinstance(visual.get("kind"), str):
        return False
    layout = _LAYOUTS.get(visual.get("kind"))
    if layout is None:
        return False
    scene = _Scene()
    try:
        layout(scene, visual)
    except _DoesNotFit:
        return False
    palette = _palette(colors)
    for kind, item in scene.items:
        if kind == "text":
            _paint_text(slide, item, font, palette)
        elif kind == "rule":
            _paint_rule(slide, item, palette)
        elif kind == "chart":
            _paint_chart(slide, item, font, palette)
        elif kind == "table":
            _paint_table(slide, item, font, palette)
    return True
